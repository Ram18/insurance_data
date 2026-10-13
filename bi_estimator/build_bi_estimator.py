"""Generate the BI Migration Estimator v2 Excel workbook.

Run:  python build_bi_estimator.py [output.xlsx] [--regression]

Every number in the workbook is a live formula over the input and reference
tabs. --regression writes a copy configured like the old calculator's base case
(used to prove the carried-over values are wired correctly).
"""
import sys
import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = next((a for a in sys.argv[1:] if not a.startswith("--")), "BI_Migration_Estimator_v2.xlsx")
REGRESSION = "--regression" in sys.argv

# ---------------------------------------------------------------- styles
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F3864")
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_SECTION = Font(name=FONT, size=12, bold=True, color="1F3864")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_KPI = Font(name=FONT, size=18, bold=True, color="1F3864")
F_KPI_LABEL = Font(name=FONT, size=9, color="595959")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_REF = PatternFill("solid", fgColor="EDEDED")
FILL_KPI = PatternFill("solid", fgColor="F2F2F2")
FILL_GOOD = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

FMT_H = '#,##0;(#,##0);"-"'
FMT_H1 = '#,##0.0;(#,##0.0);"-"'
FMT_H2 = '#,##0.00;(#,##0.00);"-"'
FMT_USD = '$#,##0;($#,##0);"-"'
FMT_USD2 = '$#,##0.00;($#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT0 = '0%;(0%);"-"'
FMT_X = '0.00"x"'
FMT_DATE = "dd-mmm-yyyy"
FMT_INT = '0;(0);"-"'

PHASES = [
    "Initiation & assessment",
    "Architecture & design",
    "Semantic model build & consolidation",
    "Report development & unit validation",
    "SIT & optimization",
    "UAT support",
    "Deployment support",
    "Hypercare",
    "Security, operations & handover",
    "AI readiness",
    "Project management",
]
NPH = len(PHASES)
NPAIRS = 20
NWEEKS = 104
NROLES = 8
NWAVES = 5

wb = Workbook()
wb.remove(wb.active)
NAMES = {}


def name(nm, ws, ref):
    """Define a workbook-level name pointing at ws!ref (absolute)."""
    if ":" in ref:
        a, b = ref.split(":")
        txt = f"'{ws.title}'!{absref(a)}:{absref(b)}"
    else:
        txt = f"'{ws.title}'!{absref(ref)}"
    wb.defined_names[nm] = DefinedName(nm, attr_text=txt)
    NAMES[nm] = txt


def absref(a):
    col = "".join(c for c in a if c.isalpha())
    row = "".join(c for c in a if c.isdigit())
    return f"${col}${row}"


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def put(ws, ref, value, font=F_BASE, fmt=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def inp(ws, ref, value, fmt=None, note=None):
    c = put(ws, ref, value, font=F_INPUT, fmt=fmt, fill=FILL_INPUT, border=BOX)
    c.protection = Protection(locked=False)
    if note:
        c.comment = Comment(note, "BI Estimator")
    return c


def calc(ws, ref, formula, fmt=None, bold=False, fill=None):
    return put(ws, ref, formula, font=F_BOLD if bold else F_BASE, fmt=fmt, fill=fill, border=BOX)


def header_row(ws, row, col0, labels, height=30):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = CENTER
        c.border = BOX
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    put(ws, "A1", text, font=F_TITLE)
    if sub:
        put(ws, "A2", sub, font=F_SUB)
    ws.sheet_view.showGridLines = False


def section(ws, ref, text, span=4):
    c = put(ws, ref, text, font=F_SECTION, fill=FILL_SECTION)
    col = ws[ref].column
    row = ws[ref].row
    for i in range(1, span):
        ws.cell(row=row, column=col + i).fill = FILL_SECTION


def dv_list(ws, rng, source, allow_free=False, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.showErrorMessage = not allow_free
    if prompt:
        dv.promptTitle = "Input"
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_num(ws, rng, lo, hi, kind="decimal", msg=None):
    dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Out of range"
    dv.error = msg or f"Enter a value between {lo} and {hi}."
    ws.add_data_validation(dv)
    dv.add(rng)


def interp(x, X, Y):
    """Piecewise-linear interpolation over 5 anchors, flat below, extrapolated above."""
    k = f"MATCH({x},{X},1)"
    return (
        f"IF({x}<=INDEX({X},1),INDEX({Y},1),"
        f"IF({x}>=INDEX({X},5),INDEX({Y},5)+({x}-INDEX({X},5))*(INDEX({Y},5)-INDEX({Y},4))/(INDEX({X},5)-INDEX({X},4)),"
        f"INDEX({Y},{k})+({x}-INDEX({X},{k}))*(INDEX({Y},{k}+1)-INDEX({Y},{k}))/(INDEX({X},{k}+1)-INDEX({X},{k}))))"
    )


# ======================================================================
# Sheets (order = tab order)
# ======================================================================
ws_guide = wb.create_sheet("Guide")
ws_sum = wb.create_sheet("Summary")
ws_set = wb.create_sheet("Setup")
ws_scope = wb.create_sheet("Scope")
ws_risk = wb.create_sheet("Risks")
ws_roles = wb.create_sheet("Roles")
ws_rates = wb.create_sheet("Rates")
ws_lists = wb.create_sheet("Lists")
ws_base = wb.create_sheet("Baseline")
ws_eff = wb.create_sheet("Effort")
ws_sch = wb.create_sheet("Schedule")
ws_plan = wb.create_sheet("Plan")
ws_cost = wb.create_sheet("Cost")
ws_asm = wb.create_sheet("Assumptions")
ws_chk = wb.create_sheet("Checks")
ws_scn = wb.create_sheet("Scenarios")
ws_log = wb.create_sheet("ChangeLog")

for ws in wb.worksheets:
    ws.sheet_properties.outlinePr.summaryBelow = False

# ======================================================================
# Lists
# ======================================================================
ws = ws_lists
title(ws, "Lists", "Reference lists behind the drop-downs. BI team maintains; add rows in the blank space - no formula changes needed.")
ARCH = ["Enterprise semantic-layer", "Paginated / pixel-perfect", "Modern visual / self-service", "Modern code-based semantic"]
TOOLS = [
    ("SAP BusinessObjects", ARCH[0]), ("OBIEE", ARCH[0]), ("Oracle Analytics Server", ARCH[0]),
    ("MicroStrategy", ARCH[0]), ("Cognos", ARCH[0]), ("SSRS", ARCH[1]), ("Crystal Reports", ARCH[1]),
    ("Tableau", ARCH[2]), ("Power BI", ARCH[2]), ("QuickSight", ARCH[2]), ("Qlik Sense", ARCH[2]),
    ("QlikView", ARCH[2]), ("Looker Studio", ARCH[2]), ("Looker", ARCH[3]),
]
TOOL_ROWS = 40
section(ws, "A4", "BI tools and archetypes", 2)
header_row(ws, 5, 1, ["BI tool", "Archetype"])
for i in range(TOOL_ROWS):
    r = 6 + i
    t, a = TOOLS[i] if i < len(TOOLS) else (None, None)
    inp(ws, f"A{r}", t)
    inp(ws, f"B{r}", a)
dv_list(ws, f"B6:B{5 + TOOL_ROWS}", "=ArchNames")
name("ToolNames", ws, f"A6:A{5 + TOOL_ROWS}")
name("ToolArch", ws, f"B6:B{5 + TOOL_ROWS}")

section(ws, "D4", "Archetypes", 1)
header_row(ws, 5, 4, ["Archetype"])
for i, a in enumerate(ARCH):
    put(ws, f"D{6 + i}", a, border=BOX, fill=FILL_REF)
name("ArchNames", ws, "D6:D9")

# archetype factor matrix  (rows = source, cols = target)
section(ws, "F4", "Pair factor fallback by archetype (source row x target column) - proposal, BI team to confirm", 5)
header_row(ws, 5, 6, ["Source \\ Target"] + ARCH, height=45)
AF = [[1.10, 1.00, 1.10, 1.20], [1.00, 0.90, 1.05, 1.20], [1.10, 1.00, 0.95, 1.10], [1.00, 1.00, 1.00, 1.00]]
for i, a in enumerate(ARCH):
    put(ws, f"F{6 + i}", a, font=F_BOLD, border=BOX)
    for j in range(4):
        inp(ws, f"{L(7 + j)}{6 + i}", AF[i][j], fmt="0.00")
name("ArchFactor", ws, "G6:J9")

section(ws, "F12", "Semantic-model approach by archetype (source row x target column)", 5)
header_row(ws, 13, 6, ["Source \\ Target"] + ARCH, height=45)
AP = [
    ["Redesign", "Redesign", "Redesign", "Author (LookML)"],
    ["Shared datasets", "Shared datasets", "Shared datasets", "Author (LookML)"],
    ["Redesign", "Shared datasets", "Translate", "Author (LookML)"],
    ["Translate", "Shared datasets", "Translate", "Translate"],
]
for i, a in enumerate(ARCH):
    put(ws, f"F{14 + i}", a, font=F_BOLD, border=BOX)
    for j in range(4):
        inp(ws, f"{L(7 + j)}{14 + i}", AP[i][j])
dv_list(ws, "G14:J17", "=ApproachNames")
name("ApproachMatrix", ws, "G14:J17")

# known pairs (seeded with old calculator X-factors)
KNOWN = [
    ("OBIEE", "Power BI", 1.00), ("Cognos", "Power BI", 1.10), ("SAP BusinessObjects", "Power BI", 1.25),
    ("Tableau", "Power BI", 0.95), ("SSRS", "Power BI", 1.05), ("OBIEE", "QuickSight", 1.20),
    ("Cognos", "QuickSight", 1.25), ("Tableau", "QuickSight", 1.10), ("SSRS", "QuickSight", 1.20),
    ("SAP BusinessObjects", "QuickSight", 1.30), ("Cognos", "Looker", 1.25), ("OBIEE", "Looker", 1.20),
    ("Tableau", "Looker", 1.10), ("SSRS", "Looker", 1.20), ("SAP BusinessObjects", "Looker", 1.30),
]
KP_ROWS = 40
section(ws, "L4", "Known pair factors (seeded from old calculator BD7:BD21)", 4)
header_row(ws, 5, 12, ["Source BI tool", "Target BI tool", "Factor", "Key"])
for i in range(KP_ROWS):
    r = 6 + i
    s, t, f = KNOWN[i] if i < len(KNOWN) else (None, None, None)
    inp(ws, f"L{r}", s)
    inp(ws, f"M{r}", t)
    inp(ws, f"N{r}", f, fmt="0.00")
    calc(ws, f"O{r}", f'=IF(L{r}="","",L{r}&"|"&M{r})')
dv_list(ws, f"L6:M{5 + KP_ROWS}", "=ToolNames", allow_free=True)
name("KnownKey", ws, f"O6:O{5 + KP_ROWS}")
name("KnownFactor", ws, f"N6:N{5 + KP_ROWS}")

DPS = ["Oracle", "SQL Server", "Teradata", "Netezza", "DB2", "SAP HANA", "Snowflake", "Databricks",
       "Google BigQuery", "Amazon Redshift", "Azure Synapse", "Microsoft Fabric", "PostgreSQL", "Other"]
section(ws, "Q4", "Data platforms", 1)
header_row(ws, 5, 17, ["Data platform"])
for i in range(25):
    inp(ws, f"Q{6 + i}", DPS[i] if i < len(DPS) else None)
name("DPList", ws, "Q6:Q30")

SMALL = {
    "YesNo": ["Yes", "No"],
    "EngTypes": ["Pure BI", "Bundled with ETL/EDW programme"],
    "SolveModes": ["Fix team, compute iterations", "Fix iterations, compute team"],
    "CostModes": ["Role mix", "Blended"],
    "IterLens": [1, 2],
}
col = 19
for nm, vals in SMALL.items():
    header_row(ws, 5, col, [nm])
    for i, v in enumerate(vals):
        put(ws, f"{L(col)}{6 + i}", v, border=BOX, fill=FILL_REF)
    name(nm, ws, f"{L(col)}6:{L(col)}{5 + len(vals)}")
    col += 1
setw(ws, {"A": 26, "B": 30, "D": 30, "F": 30, "G": 18, "H": 18, "I": 18, "J": 18, "L": 22, "M": 18, "N": 9,
          "O": 30, "Q": 20, "S": 8, "T": 30, "U": 30, "V": 12, "W": 10})

# ======================================================================
# Baseline
# ======================================================================
ws = ws_base
title(ws, "Baseline", "Reference values used by the engine. Status shows where each value came from. Change values here only with BI-team sign-off.")
header_row(ws, 4, 1, ["Parameter", "Value", "Unit", "Status", "Source / note"])
setw(ws, {"A": 46, "B": 12, "C": 18, "D": 14, "E": 70})
brow = [5]


def bline(label, value, unit, status, note, nm=None, fmt=None):
    r = brow[0]
    put(ws, f"A{r}", label, border=BOX)
    c = inp(ws, f"B{r}", value, fmt=fmt)
    put(ws, f"C{r}", unit, border=BOX)
    put(ws, f"D{r}", status, border=BOX)
    put(ws, f"E{r}", note, border=BOX, align=WRAP)
    if nm:
        name(nm, ws, f"B{r}")
    brow[0] += 1
    return r


def bsection(text):
    brow[0] += 1
    section(ws, f"A{brow[0]}", text, 5)
    brow[0] += 1


OLD = "Old calculator"
bsection("Report development")
r0 = brow[0]
for lab, v, cell in [("Simple", 3.625, "F5"), ("Medium", 6.625, "F6"), ("Complex", 9.125, "F7"), ("Very Complex", 12.625, "F8")]:
    bline(f"Hours per report page - {lab}", v, "h / page", "Carried", f"{OLD} {cell}", fmt="0.000")
name("HrsPage", ws, f"B{r0}:B{r0 + 3}")
bline("Hours per derived (similar) report page", 1, "h / page", "Carried", f"{OLD} F9", "HrsDerived", "0.000")
r0 = brow[0]
LQ = [0.19] * 4 if REGRESSION else [0, 0.10, 0.15, 0.20]
for i, v in enumerate(LQ):
    bline(f"Learning gain - quarter {i + 1} of iterations", v, "% saving", "Reshaped",
          "User decision 0/10/15/20% (old file 0/15/20/25% by page-tier phase). Applied by quarter of the iteration run.", fmt=FMT_PCT0)
name("LearnQ", ws, f"B{r0}:B{r0 + 3}")
r = brow[0]
put(ws, f"A{r}", "Average learning gain (effort is spread evenly over iterations)", border=BOX)
calc(ws, f"B{r}", "=AVERAGE(LearnQ)", FMT_PCT)
put(ws, f"D{r}", "Calculated", border=BOX)
name("LearnAvg", ws, f"B{r}")
brow[0] += 1
bline("Paginated report effort factor", 1.2, "x", "New proposal", "Paginated / pixel-perfect pages vs interactive pages", "PaginatedFactor", FMT_X)
bline("Data-platform change uplift on custom-SQL reports", 0.25, "% uplift", "New proposal", "Applied to the custom-SQL share when source and target data platforms differ", "DPRepUplift", FMT_PCT0)
bline("Merge hours per consolidated report", 2, "h / report", "New proposal", "Effort to fold a consolidated report into the report that absorbs it", "MergeHrs", FMT_H1)

bsection("Semantic models")
r0 = brow[0]
for lab, v in [("Simple", 80), ("Medium", 120), ("Complex", 200), ("Very Complex", 320)]:
    bline(f"Model build hours - {lab}", v, "h / model", "New proposal", "Replaces old 16/40/80 by reports-per-model (E16); SOP minimum 80 h", fmt=FMT_H)
name("ModelHrs", ws, f"B{r0}:B{r0 + 3}")
bline("Consolidation hours per source model merged", 16, "h / model", "New proposal", "Applied to (source models - target models)", "ConsolHrs", FMT_H)
bline("Data-platform change uplift on semantic models", 0.15, "% uplift", "New proposal", "When source and target data platforms differ", "DPModelUplift", FMT_PCT0)
r0 = brow[0]
APPROACH = [("Redesign", 1.0), ("Translate", 0.4), ("Author (LookML)", 1.0), ("Shared datasets", 0.3)]
for lab, v in APPROACH:
    put(ws, f"A{brow[0]}", lab, border=BOX)
    inp(ws, f"B{brow[0]}", v, fmt=FMT_PCT0)
    put(ws, f"C{brow[0]}", "share of build h", border=BOX)
    put(ws, f"D{brow[0]}", "New proposal", border=BOX)
    put(ws, f"E{brow[0]}", "Semantic-model approach share of full build hours", border=BOX)
    brow[0] += 1
name("ApproachNames", ws, f"A{r0}:A{r0 + 3}")
name("ApproachShares", ws, f"B{r0}:B{r0 + 3}")

bsection("Initiation and assessment (interpolation anchors on inventory pages)")
header_row(ws, brow[0], 1, ["Anchor", "Pages", "Initiation h", "Analysis h", "Source"])
brow[0] += 1
r0 = brow[0]
for px, ini, ana in [(100, 40, 80), (500, 60, 160), (1000, 80, 240), (5000, 120, 360), (10000, 160, 480)]:
    put(ws, f"A{brow[0]}", f"Anchor {brow[0] - r0 + 1}", border=BOX)
    inp(ws, f"B{brow[0]}", px, fmt=FMT_H)
    inp(ws, f"C{brow[0]}", ini, fmt=FMT_H)
    inp(ws, f"D{brow[0]}", ana, fmt=FMT_H)
    put(ws, f"E{brow[0]}", "Reshaped: old tier values E14/E15 become straight-line anchors (no cliffs)", border=BOX)
    brow[0] += 1
name("AnchorPages", ws, f"B{r0}:B{r0 + 4}")
name("AnchorInit", ws, f"C{r0}:C{r0 + 4}")
name("AnchorAnalysis", ws, f"D{r0}:D{r0 + 4}")
bline("Assessment hours per source semantic model", 4, "h / model", "New proposal", "Inventorying universes / RPD subject areas / packages / SQL datasets", "AssessPerModel", FMT_H1)

bsection("Architecture, security, operations, AI readiness")
bline("Architecture hours per distinct target BI tool", 80, "h / tool", "New proposal", "Tenant/workspace design, security, connectivity, CI/CD, performance, standards", "ArchPerTool", FMT_H)
bline("Architecture hours per migration pair", 24, "h / pair", "New proposal", "Pair-specific design decisions", "ArchPerPair", FMT_H)
r0 = brow[0]
for lab, v in [("None", 0), ("Role-based", 4), ("Row-level", 12), ("Row + object-level", 20)]:
    put(ws, f"A{brow[0]}", lab, border=BOX)
    inp(ws, f"B{brow[0]}", v, fmt=FMT_H)
    put(ws, f"C{brow[0]}", "h / target model", border=BOX)
    put(ws, f"D{brow[0]}", "New proposal", border=BOX)
    put(ws, f"E{brow[0]}", "Security model build and testing per target semantic model", border=BOX)
    brow[0] += 1
name("SecNames", ws, f"A{r0}:A{r0 + 3}")
name("SecHrs", ws, f"B{r0}:B{r0 + 3}")
bline("Hours per schedule / subscription / alert", 0.5, "h / item", "New proposal", "Recreate in target", "SchedHrs", FMT_H1)
bline("Hours per user group", 1, "h / group", "New proposal", "Groups, workspace access", "GroupHrs", FMT_H1)
bline("Training and handover per pair", 40, "h / pair", "New proposal", "Demo, documentation, knowledge transfer", "HandoverHrs", FMT_H)
r0 = brow[0]
for lab, v in [("None", 0), ("Basic", 16), ("Advanced", 40)]:
    put(ws, f"A{brow[0]}", lab, border=BOX)
    inp(ws, f"B{brow[0]}", v, fmt=FMT_H)
    put(ws, f"C{brow[0]}", "h / target model", border=BOX)
    put(ws, f"D{brow[0]}", "Placeholder", border=BOX)
    put(ws, f"E{brow[0]}", "AI readiness (descriptions, synonyms, Copilot / Q&A / Amazon Q / Gemini set-up). No baseline yet - BI team to set.", border=BOX)
    brow[0] += 1
name("AINames", ws, f"A{r0}:A{r0 + 2}")
name("AIHrs", ws, f"B{r0}:B{r0 + 2}")

bsection("Automation and deployment")
CAPS = [("Assessment", 0.6, "CapAssess"), ("Semantic model", 0.4, "CapSem"), ("Report development", 0.6, "CapRep"),
        ("SIT", 0.3, "CapSIT"), ("UAT support", 0.2, "CapUAT"), ("Deployment support", 0.6, "CapDep")]
cap_rows = {}
for lab, v, nm in CAPS:
    v = 1.0 if (REGRESSION and nm == "CapRep") else v
    cap_rows[nm] = bline(f"Automatable share cap - {lab}", v, "% of phase", "User decision",
                         "Max share of the phase automation can remove (requirements Q6)", nm, FMT_PCT0)
bline("Automation tool set-up per pair (when any automation > 0)", 40, "h / pair", "New proposal", "Accelerator configuration and runs", "AutoSetupHrs", FMT_H)
bline("Manual deployment hours per report (no pipeline)", 0.5, "h / report", "Carried", f"{OLD} E20", "ManualDepHrs", FMT_H1)

bsection("Contingency and controls")
r0 = brow[0]
for lab, v in [("Tool-assessed", 0.10), ("Customer inventory", 0.15), ("Verbal estimate", 0.25)]:
    put(ws, f"A{brow[0]}", lab, border=BOX)
    inp(ws, f"B{brow[0]}", v, fmt=FMT_PCT0)
    put(ws, f"C{brow[0]}", "contingency", border=BOX)
    put(ws, f"D{brow[0]}", "New proposal", border=BOX)
    put(ws, f"E{brow[0]}", "Replaces old flat 15% (B39). Driven by each pair's input confidence.", border=BOX)
    brow[0] += 1
name("ConfNames", ws, f"A{r0}:A{r0 + 2}")
name("ConfPct", ws, f"B{r0}:B{r0 + 2}")
bline("Pair factor override - minimum", 0.7, "x", "New proposal", "Overrides are clamped to this range", "FactorMin", FMT_X)
bline("Pair factor override - maximum", 2.0, "x", "New proposal", "", "FactorMax", FMT_X)
bline("Plan reconciliation tolerance", 0.05, "+/- %", "New proposal", "Plan hours vs estimate hours", "ReconTol", FMT_PCT0)

# ======================================================================
# Rates
# ======================================================================
ws = ws_rates
title(ws, "Rates", "Rate card (bill rates, USD per hour) as provided. Key = Role | Grade | Location is what the Roles tab selects.")
header_row(ws, 4, 1, ["Role", "Grade", "Location", "Rate (USD/h)", "Key"])
RATES = [
    ("MSE II", "G13", "India", 87), ("MSE I", "G12", "India", 82), ("Architect II", "G12", "India", 70),
    ("Architect I", "G11", "India", 65), ("Lead II", "G11", "India", 70), ("Lead II", "G10", "India", 70),
    ("Lead I", "G10", "India", 62), ("Lead I", "G9", "India", 62),
    ("Sr. Engineer II", "G10", "India", 62), ("Sr. Engineer II", "G9", "India", 62),
    ("Sr. Engineer I", "G9", "India", 57.5), ("Sr. Engineer I", "G8", "India", 57.5),
    ("Engineer II", "G8", "India", 50), ("Engineer II", "G6", "India", 50), ("Engineer I", "G6", "India", 42),
    ("Engineer I", "G4", "India", 42), ("Jr. Engineer", "G4", "India", 35), ("Architect II", "G11", "India", 70),
    ("Sr. Architect I", "G13", "India", 87), ("Sr. Architect II", "G14", "India", 105),
    ("Sr. Architect II", "G13", "India", 105), ("Architect II", "G12", "USA", 185), ("MSE I", "G12", "USA", 190),
    ("MSE II", "G13", "USA", 210),
]
RATE_ROWS = 34
for i in range(RATE_ROWS):
    r = 5 + i
    vals = RATES[i] if i < len(RATES) else (None, None, None, None)
    inp(ws, f"A{r}", vals[0])
    inp(ws, f"B{r}", vals[1])
    inp(ws, f"C{r}", vals[2])
    inp(ws, f"D{r}", vals[3], fmt=FMT_USD2)
    calc(ws, f"E{r}", f'=IF(A{r}="","",A{r}&" | "&B{r}&" | "&C{r})')
name("RateKeys", ws, f"E5:E{4 + RATE_ROWS}")
name("RateVals", ws, f"D5:D{4 + RATE_ROWS}")
name("RateLocs", ws, f"C5:C{4 + RATE_ROWS}")
put(ws, "G4", "Notes", font=F_BOLD)
put(ws, "G5", "Source: rate card supplied by the user (bill rates, USD/h). Lead I rows added from the requirements review. MSE = Manager Software Engineer.", align=WRAP)
ws.merge_cells("G5:J8")
setw(ws, {"A": 18, "B": 8, "C": 10, "D": 13, "E": 34, "G": 20, "H": 20, "I": 20, "J": 20})

# ======================================================================
# Roles (plan roles + role mix by phase)
# ======================================================================
ws = ws_roles
title(ws, "Roles", "Plan roles (each priced from one rate-card row) and the share of each phase's hours carried by each role.")
section(ws, "A4", "Plan roles", 5)
header_row(ws, 5, 1, ["#", "Plan role", "Rate-card key (Role | Grade | Location)", "Rate (USD/h)", "Location"])
ROLES = [
    ("Delivery manager / PM", "MSE I | G12 | India"),
    ("Onsite architect", "Architect II | G12 | USA"),
    ("Solution architect", "Sr. Architect I | G13 | India"),
    ("BI architect", "Architect II | G11 | India"),
    ("Lead", "Lead II | G11 | India"),
    ("Sr. engineer", "Sr. Engineer II | G10 | India"),
    ("Engineer", "Engineer II | G8 | India"),
    ("QA engineer", "Sr. Engineer I | G9 | India"),
]
for i, (rn, key) in enumerate(ROLES):
    r = 6 + i
    put(ws, f"A{r}", i + 1, border=BOX)
    inp(ws, f"B{r}", rn)
    inp(ws, f"C{r}", key)
    calc(ws, f"D{r}", f'=IFERROR(INDEX(RateVals,MATCH(C{r},RateKeys,0)),0)', FMT_USD2)
    calc(ws, f"E{r}", f'=IFERROR(INDEX(RateLocs,MATCH(C{r},RateKeys,0)),"")')
dv_list(ws, "C6:C13", "=RateKeys")
name("RoleNames", ws, "B6:B13")
name("RoleRates", ws, "D6:D13")
name("RoleLocs", ws, "E6:E13")

MIX_ROW0 = 18
section(ws, "A16", "Role mix by phase (share of phase hours; each row must total 100%)", 11)
header_row(ws, 17, 1, ["#", "Phase"] + [f"=B{6 + j}" for j in range(NROLES)] + ["Total"], height=45)
MIX = [
    [0, .20, .30, .30, .20, 0, 0, 0],
    [0, .25, .40, .35, 0, 0, 0, 0],
    [0, 0, 0, .15, .15, .50, .20, 0],
    [0, 0, 0, 0, .10, .30, .60, 0],
    [0, 0, 0, 0, .10, 0, .20, .70],
    [0, 0, 0, 0, .20, .40, .40, 0],
    [0, 0, 0, 0, .30, .70, 0, 0],
    [0, 0, 0, 0, .25, .50, .25, 0],
    [0, 0, 0, .30, .30, .40, 0, 0],
    [0, 0, 0, .40, 0, .60, 0, 0],
    [1, 0, 0, 0, 0, 0, 0, 0],
]
for p in range(NPH):
    r = MIX_ROW0 + p
    put(ws, f"A{r}", p + 1, border=BOX)
    put(ws, f"B{r}", PHASES[p], border=BOX)
    for j in range(NROLES):
        inp(ws, f"{L(3 + j)}{r}", MIX[p][j], fmt=FMT_PCT0)
    calc(ws, f"K{r}", f"=SUM(C{r}:J{r})", FMT_PCT0, bold=True)
name("RoleMix", ws, f"C{MIX_ROW0}:J{MIX_ROW0 + NPH - 1}")
name("RoleMixTotals", ws, f"K{MIX_ROW0}:K{MIX_ROW0 + NPH - 1}")
ws.conditional_formatting.add(f"K{MIX_ROW0}:K{MIX_ROW0 + NPH - 1}",
                              FormulaRule(formula=[f"ABS(K{MIX_ROW0}-1)>0.0001"], fill=FILL_BAD))
setw(ws, {"A": 5, "B": 36, "C": 34, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13, "J": 13, "K": 10})

# ======================================================================
# Setup
# ======================================================================
ws = ws_set
title(ws, "Setup", "Engagement-wide settings. Yellow cells are inputs; hover a cell for guidance.")
setw(ws, {"A": 2, "B": 52, "C": 34, "D": 16, "E": 16, "F": 60})


def srow(r, label, value, nm=None, fmt=None, note=None):
    put(ws, f"B{r}", label, border=BOX)
    c = inp(ws, f"C{r}", value, fmt=fmt, note=note)
    if nm:
        name(nm, ws, f"C{r}")
    return c


section(ws, "B4", "Engagement", 2)
srow(5, "Client", "Example Client", "Client")
srow(6, "Opportunity", "BI modernisation programme", "Opportunity")
srow(7, "Estimate version", "v2.0-draft", "EstVersion")
srow(8, "Prepared by", "", "PreparedBy")
srow(9, "Estimate date", dt.date(2026, 10, 13), "EstDate", FMT_DATE)
srow(10, "Currency", "USD", "Currency")
srow(11, "Planned start date (Monday)", dt.date(2027, 1, 4), "StartDate", FMT_DATE)

section(ws, "B13", "Delivery parameters", 2)
srow(14, "Iteration length (weeks)", 2, "IterLen", note="Development iteration = SIT/UAT cadence. 1 or 2 weeks.")
dv_list(ws, "C14", "=IterLens")
srow(15, "Productive hours per person-week", 40, "HrsWk", note="User decision: 40 productive hours (of 42.5 worked).")
srow(16, "Holiday allowance (stretches calendar dates only)", 0.04, "HolidayPct", FMT_PCT0)
srow(17, "Engagement type", "Pure BI", "EngType", note="Bundled = PM carried by the ETL/EDW programme, so PM hours are 0.")
dv_list(ws, "C17", "=EngTypes")
srow(18, "Solve mode", "Fix team, compute iterations", "SolveMode")
dv_list(ws, "C18", "=SolveModes")
srow(19, "Peak report-development team (FTE) - used when fixing team", 10, "PeakTeamIn")
srow(20, "Development iterations - used when fixing iterations", 8, "IterIn")
srow(21, "Ramp-up weeks (report development)", 3, "RampUp")
srow(22, "Ramp-down weeks (report development)", 3, "RampDown")
srow(23, "Assessment team (FTE)", 3, "AssessTeam")
srow(24, "Architecture team (FTE)", 2, "ArchTeam")
srow(25, "Architecture overlap with assessment", 0.5, "ArchOverlap", FMT_PCT0)
srow(26, "Model lead time before report iterations (weeks)", 2, "ModelLead", note="Old calculator B13 - now used.")
srow(27, "Headcount rounding step (0 = no rounding)", 0.5, "RoundStep")
srow(28, "PM FTE (pure BI engagements)", 1, "PMFTE")
for ref, lo, hi in [("C19", 0.5, 500), ("C20", 1, 200), ("C21", 0, 26), ("C22", 0, 26), ("C23", 0.5, 50),
                    ("C24", 0.5, 50), ("C25", 0, 1), ("C26", 0, 26), ("C27", 0, 1), ("C28", 0, 5), ("C15", 1, 60), ("C16", 0, 0.3)]:
    dv_num(ws, ref, lo, hi)

section(ws, "B30", "SIT, UAT and deployment support (sized per development iteration)", 2)
srow(31, "SIT weeks per iteration", 1, "SITWk")
srow(32, "SIT minimum engineers", 1, "SITMin")
srow(33, "SIT crew sizing (% of dev hours per iteration)", 0.15, "SITPct", FMT_PCT0, note="Old calculator B14 (15%) now sizes the per-iteration SIT crew.")
srow(34, "SIT fixed hours per iteration (environments, regression)", 8, "SITFixed")
srow(35, "Final regression cycle (weeks)", 1, "SITFinal")
srow(36, "UAT support weeks per iteration", 1, "UATWk")
srow(37, "UAT engineers when volume <= threshold", 1, "UATLow")
srow(38, "UAT engineers when volume > threshold", 2, "UATHigh")
srow(39, "Volume threshold (report pages per iteration)", 150, "UATThreshold")
srow(40, "Final UAT cycle (weeks)", 1, "UATFinal")
srow(41, "Deployment support weeks per iteration", 1, "DepWk", note="Usually less than the UAT support cycle.")
srow(42, "Deployment support engineers", 1, "DepEng")
srow(43, "Deployment pipeline available?", "Yes", "Pipeline", note="No = add manual deployment hours per report.")
dv_list(ws, "C43", "=YesNo")
srow(44, "Hypercare weeks per go-live", 2, "HCWk")
srow(45, "Hypercare crew (FTE)", 2, "HCCrew")
for ref, lo, hi in [("C31", 0, 2), ("C32", 0, 20), ("C33", 0, 1), ("C34", 0, 200), ("C35", 0, 8), ("C36", 0, 2),
                    ("C37", 0, 20), ("C38", 0, 20), ("C39", 1, 100000), ("C40", 0, 8), ("C41", 0, 2), ("C42", 0, 20),
                    ("C44", 0, 26), ("C45", 0, 20)]:
    dv_num(ws, ref, lo, hi)

section(ws, "B47", "Automation by phase (global; pairs may override assessment, semantic and report automation on Scope)", 4)
header_row(ws, 48, 2, ["Phase", "Automation %", "Automatable cap", "Effective saving"])
AUTO = [("Assessment", "AutoAssess", "CapAssess"), ("Semantic model", "AutoSem", "CapSem"),
        ("Report development", "AutoRep", "CapRep"), ("SIT", "AutoSIT", "CapSIT"),
        ("UAT support", "AutoUAT", "CapUAT"), ("Deployment support", "AutoDep", "CapDep")]
for i, (lab, nm, cap) in enumerate(AUTO):
    r = 49 + i
    put(ws, f"B{r}", lab, border=BOX)
    v = 0.6 if (REGRESSION and nm == "AutoRep") else 0
    inp(ws, f"C{r}", v, fmt=FMT_PCT0)
    name(nm, ws, f"C{r}")
    calc(ws, f"D{r}", f"={cap}", FMT_PCT0)
    calc(ws, f"E{r}", f"=C{r}*D{r}", FMT_PCT)
dv_num(ws, "C49:C54", 0, 1)

section(ws, "B56", "Cost", 2)
srow(57, "Cost mode", "Role mix", "CostMode")
dv_list(ws, "C57", "=CostModes")
srow(58, "Blended rate (USD/h) - used in Blended mode", 60, "BlendedRate", FMT_USD2)
srow(59, "Discount on bill rates", 0, "DiscountPct", FMT_PCT0)
srow(60, "Pass-through: tools / licences (USD)", 0, "PassTools", FMT_USD)
srow(61, "Pass-through: travel (USD)", 0, "PassTravel", FMT_USD)
srow(62, "Pass-through: other (USD)", 0, "PassOther", FMT_USD)
dv_num(ws, "C59", 0, 0.5)

section(ws, "B64", "Payment milestones (must total 100%)", 2)
MS = [("Contract signature", .10), ("Assessment sign-off", .15), ("Architecture sign-off", .10),
      ("Mid-development", .25), ("Final go-live", .30), ("Hypercare exit", .10)]
for i, (lab, p) in enumerate(MS):
    r = 65 + i
    inp(ws, f"B{r}", lab)
    inp(ws, f"C{r}", p, fmt=FMT_PCT0)
put(ws, "B71", "Total", font=F_BOLD, border=BOX)
calc(ws, "C71", "=SUM(C65:C70)", FMT_PCT0, bold=True)
name("MilestoneNames", ws, "B65:B70")
name("MilestonePct", ws, "C65:C70")
name("MilestoneTotal", ws, "C71")

# ======================================================================
# Scope
# ======================================================================
ws = ws_scope
title(ws, "Scope", "One row per migration pair (source BI tool + data platform -> target BI tool + data platform). Rows with a source, target and report count are active.")
SC = [  # key, header, width, fmt
    ("no", "#", 4, None), ("pair", "Pair name", 26, None), ("src", "Source BI tool", 18, None),
    ("tgt", "Target BI tool", 16, None), ("sdp", "Source data platform", 16, None),
    ("tdp", "Target data platform", 16, None), ("wave", "Release wave (1-5)", 9, FMT_INT),
    ("reports", "Reports in inventory", 11, FMT_H), ("ppr", "Avg pages (sheets) per report", 11, "0.0"),
    ("pag", "% paginated", 10, FMT_PCT0), ("sql", "% with custom SQL", 10, FMT_PCT0),
    ("rs", "% Simple", 9, FMT_PCT0), ("rm", "% Medium", 9, FMT_PCT0), ("rc", "% Complex", 9, FMT_PCT0),
    ("rv", "% Very Complex", 9, FMT_PCT0), ("ret", "% Retired", 9, FMT_PCT0),
    ("con", "% Consolidated", 11, FMT_PCT0), ("der", "% Derived", 9, FMT_PCT0),
    ("srcm", "Source semantic models / SQL datasets", 12, FMT_H), ("ms", "Model % Simple", 9, FMT_PCT0),
    ("mm", "Model % Medium", 9, FMT_PCT0), ("mc", "Model % Complex", 9, FMT_PCT0),
    ("mv", "Model % Very Complex", 9, FMT_PCT0), ("tgtm", "Target semantic models", 10, FMT_H),
    ("sec", "Security model", 16, None), ("sched", "Schedules / subscriptions / alerts", 12, FMT_H),
    ("groups", "User groups", 9, FMT_H), ("ai", "AI readiness", 11, None),
    ("fov", "Pair factor override", 10, FMT_X), ("why", "Override reason", 28, None),
    ("aa", "Automation override - Assessment", 11, FMT_PCT0), ("as", "Automation override - Semantic", 11, FMT_PCT0),
    ("ar", "Automation override - Reports", 11, FMT_PCT0), ("conf", "Input confidence", 17, None),
    ("notes", "Notes", 30, None),
]
SCOL = {k: L(i + 1) for i, (k, *_r) in enumerate(SC)}
SR0, SR1 = 7, 6 + NPAIRS
header_row(ws, 6, 1, [h for _k, h, *_r in SC], height=60)
for i, (k, h, w, f) in enumerate(SC):
    ws.column_dimensions[L(i + 1)].width = w
groups = [("Pair", "A", "G"), ("Report inventory", "H", "K"), ("Report complexity", "L", "O"),
          ("Rationalisation", "P", "R"), ("Semantic layer", "S", "X"), ("Security, ops, AI", "Y", "AB"),
          ("Overrides and confidence", "AC", "AI")]
for gname, a, b in groups:
    ws.merge_cells(f"{a}5:{b}5")
    put(ws, f"{a}5", gname, font=F_BOLD, fill=FILL_SECTION, align=CENTER)
EX = [
    dict(pair="Finance OBIEE to Power BI", src="OBIEE", tgt="Power BI", sdp="Oracle", tdp="Snowflake", wave=1,
         reports=400, ppr=3, pag=0.10, sql=0.15, rs=.40, rm=.35, rc=.20, rv=.05, ret=.20, con=.10, der=.10,
         srcm=12, ms=.30, mm=.40, mc=.20, mv=.10, tgtm=8, sec="Row-level", sched=150, groups=20, ai="Basic",
         conf="Customer inventory", notes="Example row - replace with real scope"),
    dict(pair="Sales Tableau to Power BI", src="Tableau", tgt="Power BI", sdp="Amazon Redshift", tdp="Snowflake",
         wave=2, reports=250, ppr=2, pag=0, sql=0.10, rs=.50, rm=.30, rc=.15, rv=.05, ret=.25, con=.05, der=.10,
         srcm=40, ms=.50, mm=.30, mc=.15, mv=.05, tgtm=10, sec="Role-based", sched=60, groups=10, ai="None",
         conf="Verbal estimate", notes="Example row - replace with real scope"),
]
if REGRESSION:
    EX = [dict(pair="Regression: old base case", src="Tableau", tgt="QuickSight", sdp="", tdp="", wave=1,
               reports=5000, ppr=1, pag=0, sql=0, rs=.4, rm=.4, rc=.2, rv=0, ret=0, con=0, der=0, srcm=1000,
               ms=1, mm=0, mc=0, mv=0, tgtm=1000, sec="None", sched=0, groups=0, ai="None", conf="Customer inventory",
               notes="Old calculator base case")]
for i in range(NPAIRS):
    r = SR0 + i
    put(ws, f"A{r}", i + 1, border=BOX)
    ex = EX[i] if i < len(EX) else {}
    for k, h, w, f in SC[1:]:
        inp(ws, f"{SCOL[k]}{r}", ex.get(k), fmt=f)
rng = lambda k: f"{SCOL[k]}{SR0}:{SCOL[k]}{SR1}"
dv_list(ws, rng("src"), "=ToolNames", allow_free=True, prompt="Pick or type a BI tool. A new tool must be added on Lists with its archetype.")
dv_list(ws, rng("tgt"), "=ToolNames", allow_free=True)
dv_list(ws, rng("sdp"), "=DPList", allow_free=True)
dv_list(ws, rng("tdp"), "=DPList", allow_free=True)
dv_list(ws, rng("sec"), "=SecNames")
dv_list(ws, rng("ai"), "=AINames")
dv_list(ws, rng("conf"), "=ConfNames")
dv_num(ws, rng("wave"), 1, NWAVES, "whole")
for k in ("pag", "sql", "rs", "rm", "rc", "rv", "ret", "con", "der", "ms", "mm", "mc", "mv", "aa", "as", "ar"):
    dv_num(ws, rng(k), 0, 1)
dv_num(ws, rng("ppr"), 1, 100, msg="Pages per report must be at least 1.")
for k in ("reports", "srcm", "tgtm", "sched", "groups"):
    dv_num(ws, rng(k), 0, 1000000)
dv_num(ws, rng("fov"), 0.7, 2.0)
ws.freeze_panes = "C7"
put(ws, "A3", "Mixes are fractions that total 100% per row. Retired + Consolidated <= 100%; Derived applies to what remains. Blank automation override = use Setup.", font=F_SUB)


def S(k, r):
    return f"Scope!${SCOL[k]}{r}"


# ======================================================================
# Effort (per pair + programme)
# ======================================================================
ws = ws_eff
title(ws, "Effort", "Hours per pair and per phase. Locked - every value is calculated from Scope, Setup and Baseline.")
EC = [  # key, header, fmt, formula(r)
    ("no", "#", None, lambda r: f"={S('no', r)}"),
    ("pair", "Pair", None, lambda r: f'=IF({S("pair", r)}="","",{S("pair", r)})'),
    ("act", "Active", FMT_INT, lambda r: f'=IF(AND({S("src", r)}<>"",{S("tgt", r)}<>"",N({S("reports", r)})>0),1,0)'),
    ("sarch", "Source archetype", None, lambda r: f'=IF(C{r}=0,"",IFERROR(INDEX(ToolArch,MATCH({S("src", r)},ToolNames,0)),"Unmapped"))'),
    ("tarch", "Target archetype", None, lambda r: f'=IF(C{r}=0,"",IFERROR(INDEX(ToolArch,MATCH({S("tgt", r)},ToolNames,0)),"Unmapped"))'),
    ("fsug", "Suggested pair factor", FMT_X, lambda r: f'=IF(C{r}=0,0,IFERROR(INDEX(KnownFactor,MATCH({S("src", r)}&"|"&{S("tgt", r)},KnownKey,0)),IFERROR(INDEX(ArchFactor,MATCH(D{r},ArchNames,0),MATCH(E{r},ArchNames,0)),1)))'),
    ("f", "Pair factor used", FMT_X, lambda r: f'=IF(C{r}=0,0,IF(N({S("fov", r)})>0,MIN(FactorMax,MAX(FactorMin,{S("fov", r)})),F{r}))'),
    ("appr", "Semantic approach", None, lambda r: f'=IF(C{r}=0,"",IFERROR(INDEX(ApproachMatrix,MATCH(D{r},ArchNames,0),MATCH(E{r},ArchNames,0)),"Redesign"))'),
    ("ashare", "Approach share", FMT_PCT0, lambda r: f'=IF(C{r}=0,0,IFERROR(INDEX(ApproachShares,MATCH(H{r},ApproachNames,0)),1))'),
    ("dpchg", "Data platform change", FMT_INT, lambda r: f'=IF(AND(C{r}=1,{S("sdp", r)}<>"",{S("tdp", r)}<>"",{S("sdp", r)}<>{S("tdp", r)}),1,0)'),
    ("urep", "Report platform uplift", FMT_X, lambda r: f'=1+J{r}*N({S("sql", r)})*DPRepUplift'),
    ("umod", "Model platform uplift", FMT_X, lambda r: f'=1+J{r}*DPModelUplift'),
    ("ttype", "Report-type factor", FMT_X, lambda r: f'=1+N({S("pag", r)})*(PaginatedFactor-1)'),
    ("svas", "Saving - assessment", FMT_PCT, lambda r: f'=IF({S("aa", r)}="",AutoAssess,{S("aa", r)})*CapAssess'),
    ("svsm", "Saving - semantic", FMT_PCT, lambda r: f'=IF({S("as", r)}="",AutoSem,{S("as", r)})*CapSem'),
    ("svrp", "Saving - reports", FMT_PCT, lambda r: f'=IF({S("ar", r)}="",AutoRep,{S("ar", r)})*CapRep'),
    ("invp", "Inventory pages", FMT_H, lambda r: f'=C{r}*N({S("reports", r)})*N({S("ppr", r)})'),
    ("build", "Reports to build", FMT_H1, lambda r: f'=C{r}*N({S("reports", r)})*MAX(0,1-N({S("ret", r)})-N({S("con", r)}))'),
    ("derr", "Derived reports", FMT_H1, lambda r: f'=R{r}*N({S("der", r)})'),
    ("fullr", "Full-effort reports", FMT_H1, lambda r: f'=R{r}-S{r}'),
    ("fullp", "Full-effort pages", FMT_H1, lambda r: f'=T{r}*N({S("ppr", r)})'),
    ("derp", "Derived pages", FMT_H1, lambda r: f'=S{r}*N({S("ppr", r)})'),
    ("conr", "Consolidated reports", FMT_H1, lambda r: f'=C{r}*N({S("reports", r)})*N({S("con", r)})'),
    ("hpp", "Mix-weighted h / page", FMT_H2, lambda r: f'=N({S("rs", r)})*INDEX(HrsPage,1)+N({S("rm", r)})*INDEX(HrsPage,2)+N({S("rc", r)})*INDEX(HrsPage,3)+N({S("rv", r)})*INDEX(HrsPage,4)'),
    ("repbase", "Report base hours", FMT_H, lambda r: f'=(U{r}*X{r}*K{r}*M{r}+V{r}*HrsDerived)*G{r}'),
    ("merge", "Merge hours", FMT_H, lambda r: f'=W{r}*MergeHrs*G{r}'),
    ("rep", "REPORT DEVELOPMENT h", FMT_H, lambda r: f'=Y{r}*(1-LearnAvg)*(1-P{r})+Z{r}'),
    ("tm", "Target models", FMT_H1, lambda r: f'=C{r}*N({S("tgtm", r)})'),
    ("hpm", "Mix-weighted h / model", FMT_H1, lambda r: f'=N({S("ms", r)})*INDEX(ModelHrs,1)+N({S("mm", r)})*INDEX(ModelHrs,2)+N({S("mc", r)})*INDEX(ModelHrs,3)+N({S("mv", r)})*INDEX(ModelHrs,4)'),
    ("mbuild", "Model build hours", FMT_H, lambda r: f'=AB{r}*AC{r}*I{r}*G{r}*L{r}'),
    ("consol", "Consolidation hours", FMT_H, lambda r: f'=C{r}*MAX(0,N({S("srcm", r)})-N({S("tgtm", r)}))*ConsolHrs*G{r}'),
    ("sem", "SEMANTIC MODEL h", FMT_H, lambda r: f'=(AD{r}+AE{r})*(1-O{r})'),
    ("ana", "ASSESSMENT (analysis) h", FMT_H, lambda r: f'=IF(C{r}=0,0,({interp(f"Q{r}", "AnchorPages", "AnchorAnalysis")}*G{r}+N({S("srcm", r)})*AssessPerModel)*(1-N{r}))'),
    ("secops", "SECURITY, OPS & HANDOVER h", FMT_H, lambda r: f'=IF(C{r}=0,0,AB{r}*IFERROR(INDEX(SecHrs,MATCH({S("sec", r)},SecNames,0)),0)+N({S("sched", r)})*SchedHrs+N({S("groups", r)})*GroupHrs+HandoverHrs)'),
    ("aih", "AI READINESS h", FMT_H, lambda r: f'=AB{r}*IFERROR(INDEX(AIHrs,MATCH({S("ai", r)},AINames,0)),0)'),
    ("autoset", "Automation set-up h", FMT_H, lambda r: f'=IF(AND(C{r}=1,MAX(N{r},O{r},P{r},AutoSIT,AutoUAT,AutoDep)>0),AutoSetupHrs,0)'),
    ("direct", "Pair direct hours", FMT_H, lambda r: f'=AA{r}+AF{r}+AG{r}+AH{r}+AI{r}+AJ{r}'),
    ("dev", "Development hours (semantic + reports)", FMT_H, lambda r: f'=AA{r}+AF{r}'),
    ("wave", "Wave", FMT_INT, lambda r: f'=IF(C{r}=0,0,MIN({NWAVES},MAX(1,INT(N({S("wave", r)})))))'),
    ("cont", "Contingency %", FMT_PCT0, lambda r: f'=IF(C{r}=0,0,IFERROR(INDEX(ConfPct,MATCH({S("conf", r)},ConfNames,0)),MAX(ConfPct)))'),
    ("alloc", "Allocated programme hours", FMT_H, lambda r: f'=IF(SUM($AL${SR0}:$AL${SR1})>0,AL{r}/SUM($AL${SR0}:$AL${SR1})*ProgSharedHrs,0)'),
    ("total", "PAIR TOTAL h (before contingency)", FMT_H, lambda r: f'=AK{r}+AO{r}'),
    ("conth", "Contingency h", FMT_H, lambda r: f'=AP{r}*AN{r}'),
    ("fmix", "Flag: mix not 100%", FMT_INT, lambda r: f'=IF(C{r}=0,0,IF(OR(ABS(N({S("rs", r)})+N({S("rm", r)})+N({S("rc", r)})+N({S("rv", r)})-1)>0.001,AND(N({S("tgtm", r)})>0,ABS(N({S("ms", r)})+N({S("mm", r)})+N({S("mc", r)})+N({S("mv", r)})-1)>0.001)),1,0))'),
    ("frat", "Flag: rationalisation > 100%", FMT_INT, lambda r: f'=IF(AND(C{r}=1,N({S("ret", r)})+N({S("con", r)})>1),1,0)'),
    ("fmap", "Flag: unmapped tool", FMT_INT, lambda r: f'=IF(AND(C{r}=1,OR(D{r}="Unmapped",E{r}="Unmapped")),1,0)'),
    ("fwhy", "Flag: override without reason", FMT_INT, lambda r: f'=IF(AND(C{r}=1,N({S("fov", r)})>0,ROUND(N({S("fov", r)})-F{r},4)<>0,{S("why", r)}=""),1,0)'),
    ("fppr", "Flag: pages < 1 or no target model", FMT_INT, lambda r: f'=IF(AND(C{r}=1,OR(N({S("ppr", r)})<1,N({S("tgtm", r)})<1)),1,0)'),
]
ECOL = {k: L(i + 1) for i, (k, *_r) in enumerate(EC)}
# sanity: the hard-coded column letters used in formulas above must match
expect = {"act": "C", "sarch": "D", "tarch": "E", "fsug": "F", "f": "G", "appr": "H", "ashare": "I", "dpchg": "J",
          "urep": "K", "umod": "L", "ttype": "M", "svas": "N", "svsm": "O", "svrp": "P", "invp": "Q", "build": "R",
          "derr": "S", "fullr": "T", "fullp": "U", "derp": "V", "conr": "W", "hpp": "X", "repbase": "Y", "merge": "Z",
          "rep": "AA", "tm": "AB", "hpm": "AC", "mbuild": "AD", "consol": "AE", "sem": "AF", "ana": "AG",
          "secops": "AH", "aih": "AI", "autoset": "AJ", "direct": "AK", "dev": "AL", "wave": "AM", "cont": "AN",
          "alloc": "AO", "total": "AP", "conth": "AQ"}
for k, v in expect.items():
    assert ECOL[k] == v, (k, ECOL[k], v)
header_row(ws, 6, 1, [h for _k, h, *_r in EC], height=60)
for i, (k, h, f, fn) in enumerate(EC):
    ws.column_dimensions[L(i + 1)].width = 12
    for p in range(NPAIRS):
        r = SR0 + p
        calc(ws, f"{L(i + 1)}{r}", fn(r), f)
ws.column_dimensions["B"].width = 26
for k in ("sarch", "tarch", "appr"):
    ws.column_dimensions[ECOL[k]].width = 22
ER1 = SR1 + 1
put(ws, f"B{ER1}", "Total", font=F_BOLD, border=BOX)
for k in ("act", "invp", "build", "derr", "fullp", "derp", "conr", "rep", "tm", "sem", "ana", "secops", "aih",
          "autoset", "direct", "dev", "alloc", "total", "conth", "fmix", "frat", "fmap", "fwhy", "fppr"):
    c = ECOL[k]
    calc(ws, f"{c}{ER1}", f"=SUM({c}{SR0}:{c}{SR1})", FMT_H, bold=True)
for k, nm in [("act", "ActivePairs"), ("invp", "InvPages"), ("build", "ReportsBuilt"), ("rep", "RepHrs"),
              ("sem", "SemHrs"), ("ana", "AnaHrs"), ("secops", "SecOpsHrs"), ("aih", "AIHrsTotal"),
              ("autoset", "AutoSetTotal"), ("direct", "PairDirectHrs"), ("dev", "DevHrsTotal"),
              ("total", "PairTotalHrs"), ("conth", "ContHrsPairs"), ("fmix", "FlagMix"), ("frat", "FlagRat"),
              ("fmap", "FlagMap"), ("fwhy", "FlagWhy"), ("fppr", "FlagPpr")]:
    name(nm, ws, f"{ECOL[k]}{ER1}")
name("PagesFull", ws, f"{ECOL['fullp']}{ER1}")
name("PagesDerived", ws, f"{ECOL['derp']}{ER1}")
name("EffWave", ws, f"{ECOL['wave']}{SR0}:{ECOL['wave']}{SR1}")
name("EffDev", ws, f"{ECOL['dev']}{SR0}:{ECOL['dev']}{SR1}")
name("EffTotal", ws, f"{ECOL['total']}{SR0}:{ECOL['total']}{SR1}")
name("EffCont", ws, f"{ECOL['cont']}{SR0}:{ECOL['cont']}{SR1}")
ws.freeze_panes = "C7"

# programme-level block
P0 = ER1 + 3
section(ws, f"A{P0}", "Programme-level effort (shared across pairs)", 6)
header_row(ws, P0 + 1, 2, ["Item", "Hours", "Basis"])
PROG = [
    ("ProgInitHrs", "Initiation (programme)", f"={interp('InvPages', 'AnchorPages', 'AnchorInit')}*(1-AutoAssess*CapAssess)", "Interpolated on total inventory pages; old E14 anchors"),
    ("TargetTools", "Distinct target BI tools", f'=SUMPRODUCT((Scope!${SCOL["tgt"]}${SR0}:${SCOL["tgt"]}${SR1}<>"")/COUNTIF(Scope!${SCOL["tgt"]}${SR0}:${SCOL["tgt"]}${SR1},Scope!${SCOL["tgt"]}${SR0}:${SCOL["tgt"]}${SR1}&""))', "Count"),
    ("ProgArchHrs", "Architecture & design", "=TargetTools*ArchPerTool+ActivePairs*ArchPerPair", "Per target tool + per pair"),
    ("ProgSITHrs", "SIT & optimization", "=((IterCount*SITWk+SITFinal)*SITEng*HrsWk+IterCount*SITFixed)*(1-AutoSIT*CapSIT)", "Per iteration crew x weeks + fixed + final regression"),
    ("ProgUATHrs", "UAT support", "=(IterCount*UATWk+UATFinal)*UATEng*HrsWk*(1-AutoUAT*CapUAT)", "Per iteration: 1-2 engineers x 1-2 weeks by volume"),
    ("ProgDepHrs", "Deployment support", '=(IterCount*DepWk*DepEng*HrsWk+IF(Pipeline="No",ReportsBuilt*ManualDepHrs,0))*(1-AutoDep*CapDep)', "Per iteration, plus manual deployment without a pipeline"),
    ("ProgHCHrs", "Hypercare", "=Releases*HCWk*HCCrew*HrsWk", "Per go-live"),
    ("ProgPMHrs", "Project management", '=IF(EngType="Pure BI",PMFTE*DurationWeeks*HrsWk,0)', "PM FTE x duration; 0 when bundled"),
]
for i, (nm, lab, fml, basis) in enumerate(PROG):
    r = P0 + 2 + i
    put(ws, f"B{r}", lab, border=BOX)
    calc(ws, f"C{r}", fml, FMT_H if nm != "TargetTools" else FMT_INT)
    put(ws, f"D{r}", basis, border=BOX)
    name(nm, ws, f"C{r}")
r = P0 + 2 + len(PROG)
put(ws, f"B{r}", "Shared hours allocated to pairs by development hours", font=F_BOLD, border=BOX)
calc(ws, f"C{r}", "=ProgInitHrs+ProgArchHrs+ProgSITHrs+ProgUATHrs+ProgDepHrs+ProgHCHrs+ProgPMHrs", FMT_H, bold=True)
name("ProgSharedHrs", ws, f"C{r}")
ws.column_dimensions["D"].width = 22

PH0 = r + 3
section(ws, f"A{PH0}", "Hours by phase", 6)
header_row(ws, PH0 + 1, 1, ["#", "Phase", "Hours", "% of total"])
PHF = [
    "=ProgInitHrs+AnaHrs+AutoSetTotal", "=ProgArchHrs", "=SemHrs", "=RepHrs", "=ProgSITHrs", "=ProgUATHrs",
    "=ProgDepHrs", "=ProgHCHrs", "=SecOpsHrs", "=AIHrsTotal", "=ProgPMHrs",
]
for p in range(NPH):
    r = PH0 + 2 + p
    put(ws, f"A{r}", p + 1, border=BOX)
    put(ws, f"B{r}", PHASES[p], border=BOX)
    calc(ws, f"C{r}", PHF[p], FMT_H)
    calc(ws, f"D{r}", f"=IF(TotalBaseHrs>0,C{r}/TotalBaseHrs,0)", FMT_PCT)
name("PhaseNames", ws, f"B{PH0 + 2}:B{PH0 + 1 + NPH}")
name("PhaseHours", ws, f"C{PH0 + 2}:C{PH0 + 1 + NPH}")
r = PH0 + 2 + NPH
TOT = [
    ("TotalBaseHrs", "Total base hours", "=SUM(PhaseHours)", FMT_H),
    ("ContPct", "Contingency % (weighted by pair input confidence)", "=IF(SUM(EffTotal)>0,SUMPRODUCT(EffTotal,EffCont)/SUM(EffTotal),0)", FMT_PCT),
    ("ContHrs", "Contingency hours", "=TotalBaseHrs*ContPct", FMT_H),
    ("RiskHrs", "Risk reserve hours (Risks tab)", "=RiskExpectedHrs", FMT_H),
    ("TotalHrs", "TOTAL HOURS incl. contingency and risk", "=TotalBaseHrs+ContHrs+RiskHrs", FMT_H),
]
for i, (nm, lab, fml, f) in enumerate(TOT):
    put(ws, f"B{r + i}", lab, font=F_BOLD, border=BOX)
    calc(ws, f"C{r + i}", fml, f, bold=True)
    name(nm, ws, f"C{r + i}")

# ======================================================================
# Risks
# ======================================================================
ws = ws_risk
title(ws, "Risks", "Named risks priced at expected value (impact hours x probability). Example rows - replace with the engagement's risks.")
header_row(ws, 5, 1, ["#", "Risk", "Impact (hours)", "Probability", "Expected hours", "Owner", "Mitigation"])
RISKS = [("Target data platform tables not ready when report iterations start", 200, 0.3, "Customer", "Model-lead gate; mock data for unit tests"),
         ("Custom visuals / extensions with no target equivalent", 120, 0.25, "BI lead", "Identify in assessment; agree alternatives")]
for i in range(15):
    r = 6 + i
    put(ws, f"A{r}", i + 1, border=BOX)
    ex = RISKS[i] if (i < len(RISKS) and not REGRESSION) else (None, None, None, None, None)
    inp(ws, f"B{r}", ex[0])
    inp(ws, f"C{r}", ex[1], fmt=FMT_H)
    inp(ws, f"D{r}", ex[2], fmt=FMT_PCT0)
    calc(ws, f"E{r}", f"=N(C{r})*N(D{r})", FMT_H)
    inp(ws, f"F{r}", ex[3])
    inp(ws, f"G{r}", ex[4])
dv_num(ws, "D6:D20", 0, 1)
put(ws, "D21", "Total", font=F_BOLD, border=BOX)
calc(ws, "E21", "=SUM(E6:E20)", FMT_H, bold=True)
name("RiskExpectedHrs", ws, "E21")
setw(ws, {"A": 4, "B": 60, "C": 13, "D": 12, "E": 13, "F": 16, "G": 50})

# ======================================================================
# Schedule
# ======================================================================
ws = ws_sch
title(ws, "Schedule", "Iterations, windows and milestones (working weeks from the planned start). Locked.")
setw(ws, {"A": 2, "B": 56, "C": 16, "D": 16, "E": 16, "F": 16, "G": 18, "H": 20})
section(ws, "B4", "Development iterations", 3)
SCH = [
    ("AssessLen", "Assessment weeks", "=MAX(1,ROUNDUP(INDEX(PhaseHours,1)/(AssessTeam*HrsWk),0))", FMT_INT),
    ("AssessEnd", "Assessment end week", "=AssessLen", FMT_INT),
    ("ArchStart", "Architecture start week", "=1+ROUND(AssessLen*(1-ArchOverlap),0)", FMT_INT),
    ("ArchEnd", "Architecture end week", "=ArchStart+MAX(1,ROUNDUP(INDEX(PhaseHours,2)/(ArchTeam*HrsWk),0))-1", FMT_INT),
    ("DevStart", "Development start week (semantic models)", "=MAX(AssessEnd,ArchEnd)+1", FMT_INT),
    ("RepStart", "First report iteration starts (week)", "=DevStart+ModelLead", FMT_INT),
    ("IterCount", "Development iterations", '=IF(SolveMode="Fix iterations, compute team",MAX(1,ROUND(IterIn,0)),MAX(1,ROUNDUP((RepHrs/(PeakTeamIn*HrsWk)+(RampUp+RampDown)/2)/IterLen,0)))', FMT_INT),
    ("RepWeeks", "Report development weeks", "=IterCount*IterLen", FMT_INT),
    ("RepEnd", "Last report iteration ends (week)", "=RepStart+RepWeeks-1", FMT_INT),
    ("SemEnd", "Semantic model build ends (week)", "=MAX(DevStart,RepStart+ROUNDUP(RepWeeks/2,0)-1)", FMT_INT),
    ("PeakRepTeam", "Peak report-development team (FTE, before rounding)", "=IF(RampWeightSum>0,RepHrs/(HrsWk*RampWeightSum),0)", FMT_H1),
    ("DevHrsPerIter", "Development hours per iteration", "=DevHrsTotal/IterCount", FMT_H),
    ("PagesPerIter", "Report pages delivered per iteration", "=(PagesFull+PagesDerived)/IterCount", FMT_H),
    ("SITEng", "SIT engineers per iteration", "=MAX(SITMin,ROUNDUP(SITPct*DevHrsPerIter/(MAX(SITWk,0.5)*HrsWk),0))", FMT_INT),
    ("UATEng", "UAT support engineers per iteration", "=IF(PagesPerIter<=UATThreshold,UATLow,UATHigh)", FMT_INT),
    ("SITStart", "SIT window starts (one iteration behind)", "=RepStart+IterLen", FMT_INT),
    ("SITEnd", "SIT window ends (incl. final regression)", "=RepEnd+IterLen+SITFinal", FMT_INT),
    ("UATStart", "UAT window starts (two iterations behind)", "=RepStart+2*IterLen", FMT_INT),
    ("LastGoLive", "Final go-live week", "=RepEnd+2*IterLen+SITFinal+UATFinal", FMT_INT),
]
for i, (nm, lab, fml, f) in enumerate(SCH):
    r = 5 + i
    put(ws, f"B{r}", lab, border=BOX)
    calc(ws, f"C{r}", fml, f)
    name(nm, ws, f"C{r}")
r_w = 5 + len(SCH) + 2
section(ws, f"B{r_w}", "Release waves", 6)
header_row(ws, r_w + 1, 2, ["Wave", "Development hours", "Cumulative hours", "Cumulative iterations", "Has scope (wave no.)", "Go-live week"])
W0 = r_w + 2
for w in range(NWAVES):
    r = W0 + w
    put(ws, f"B{r}", w + 1, border=BOX)
    calc(ws, f"C{r}", f"=SUMIFS(EffDev,EffWave,B{r})", FMT_H)
    calc(ws, f"D{r}", f"=SUM($C${W0}:C{r})", FMT_H)
    calc(ws, f"E{r}", f"=IF(DevHrsTotal>0,ROUND(IterCount*D{r}/DevHrsTotal,0),0)", FMT_INT)
    calc(ws, f"F{r}", f"=IF(C{r}>0,B{r},0)", FMT_INT)
    calc(ws, f"G{r}", f'=IF(C{r}>0,IF(B{r}=LastWave,LastGoLive,RepStart-1+E{r}*IterLen+2*IterLen),"")', FMT_INT)
name("LastWave", ws, f"C{W0 + NWAVES + 1}")
name("WaveGoLive", ws, f"G{W0}:G{W0 + NWAVES - 1}")
r = W0 + NWAVES + 1
WV = [
    ("LastWave", "Last wave with scope", f"=MAX(F{W0}:F{W0 + NWAVES - 1})", FMT_INT),
    ("Releases", "Releases (go-lives)", f'=COUNTIF(C{W0}:C{W0 + NWAVES - 1},">0")', FMT_INT),
    ("FirstGoLive", "First go-live week", f"=IF(Releases>0,MIN(WaveGoLive),LastGoLive)", FMT_INT),
    ("HCStart", "Hypercare starts (week)", "=FirstGoLive+1", FMT_INT),
    ("ProgEnd", "Programme end week (last go-live + hypercare)", "=LastGoLive+HCWk", FMT_INT),
    ("DurationWeeks", "DURATION (working weeks)", "=ProgEnd", FMT_INT),
    ("CalendarWeeks", "Duration in calendar weeks (incl. holiday allowance)", "=ROUNDUP(DurationWeeks*(1+HolidayPct),0)", FMT_INT),
    ("EndDate", "Planned end date", "=StartDate+CalendarWeeks*7-3", FMT_DATE),
]
for i, (nm, lab, fml, f) in enumerate(WV):
    rr = r + i
    put(ws, f"B{rr}", lab, border=BOX, font=F_BOLD if nm == "DurationWeeks" else F_BASE)
    calc(ws, f"C{rr}", fml, f, bold=nm == "DurationWeeks")
    name(nm, ws, f"C{rr}")

r_p = r + len(WV) + 2
section(ws, f"B{r_p}", "Phase windows (feed the weekly Plan)", 6)
header_row(ws, r_p + 1, 2, ["Phase", "Start week", "End week", "Profile", "Hours", "Weeks"])
PW = [("1", "AssessEnd", "Flat"), ("ArchStart", "ArchEnd", "Flat"), ("DevStart", "SemEnd", "Flat"),
      ("RepStart", "RepEnd", "Ramp"), ("SITStart", "SITEnd", "Flat"), ("UATStart", "LastGoLive", "Flat"),
      ("UATStart", "LastGoLive", "Flat"), ("HCStart", "ProgEnd", "Flat"), ("RepStart", "LastGoLive", "Flat"),
      ("RepStart", "RepEnd", "Flat"), ("1", "ProgEnd", "Flat")]
PW0 = r_p + 2
for p in range(NPH):
    rr = PW0 + p
    put(ws, f"B{rr}", PHASES[p], border=BOX)
    calc(ws, f"C{rr}", f"={PW[p][0]}", FMT_INT)
    calc(ws, f"D{rr}", f"={PW[p][1]}", FMT_INT)
    put(ws, f"E{rr}", PW[p][2], border=BOX)
    calc(ws, f"F{rr}", f"=INDEX(PhaseHours,{p + 1})", FMT_H)
    calc(ws, f"G{rr}", f"=MAX(0,D{rr}-C{rr}+1)", FMT_INT)
name("PhStart", ws, f"C{PW0}:C{PW0 + NPH - 1}")
name("PhEnd", ws, f"D{PW0}:D{PW0 + NPH - 1}")

r_m = PW0 + NPH + 2
section(ws, f"B{r_m}", "Milestones", 3)
header_row(ws, r_m + 1, 2, ["Milestone", "Week", "Date (end of week)"])
MIL = [("Assessment sign-off", "=AssessEnd"), ("Architecture sign-off", "=ArchEnd"),
       ("First semantic models ready", "=RepStart-1"), ("First report iteration starts", "=RepStart"),
       ("Last report iteration ends", "=RepEnd")]
MIL += [(f"Wave {w + 1} go-live", f"=G{W0 + w}") for w in range(NWAVES)]
MIL += [("Hypercare exit / programme end", "=ProgEnd")]
M0 = r_m + 2
for i, (lab, fml) in enumerate(MIL):
    rr = M0 + i
    put(ws, f"B{rr}", lab, border=BOX)
    calc(ws, f"C{rr}", fml, FMT_INT)
    calc(ws, f"D{rr}", f'=IF(ISNUMBER(C{rr}),StartDate+ROUNDUP(C{rr}*(1+HolidayPct),0)*7-3,"")', FMT_DATE)
name("MilestoneLabels", ws, f"B{M0}:B{M0 + len(MIL) - 1}")
name("MilestoneWeeks", ws, f"C{M0}:C{M0 + len(MIL) - 1}")
name("MilestoneDates", ws, f"D{M0}:D{M0 + len(MIL) - 1}")
MIL_N = len(MIL)

# ======================================================================
# Plan (weekly resource loading)
# ======================================================================
ws = ws_plan
title(ws, "Plan", f"Weekly resource loading for up to {NWEEKS} working weeks. Phase hours are spread over their windows; report development ramps up, peaks and ramps down.")
WC0 = 6  # first week column (F)
WCL = WC0 + NWEEKS - 1
wl = lambda i: L(WC0 + i)
put(ws, "A5", "Working week", font=F_BOLD)
put(ws, "A6", "Week starting", font=F_BOLD)
header_row(ws, 4, 1, ["Row", "Start wk", "End wk", "Hours", "Profile / label"])
for i in range(NWEEKS):
    c = wl(i)
    put(ws, f"{c}5", i + 1, font=F_BOLD, align=CENTER)
    calc(ws, f"{c}6", f"=StartDate+({c}5-1)*7", "dd-mmm")
    ws.column_dimensions[c].width = 7
setw(ws, {"A": 34, "B": 8, "C": 8, "D": 10, "E": 14})
PR0 = 8
put(ws, f"A{PR0 - 1}", "Phase hours per week", font=F_SECTION)
RW = PR0 + NPH + 1  # ramp weight helper row
for p in range(NPH):
    r = PR0 + p
    put(ws, f"A{r}", PHASES[p], border=BOX)
    calc(ws, f"B{r}", f"=INDEX(PhStart,{p + 1})", FMT_INT)
    calc(ws, f"C{r}", f"=INDEX(PhEnd,{p + 1})", FMT_INT)
    calc(ws, f"D{r}", f"=INDEX(PhaseHours,{p + 1})", FMT_H)
    calc(ws, f"E{r}", f"=Schedule!E{PW0 + p}")
    for i in range(NWEEKS):
        c = wl(i)
        calc(ws, f"{c}{r}",
             f'=IF(AND({c}$5>=$B{r},{c}$5<=$C{r},$C{r}>=$B{r}),IF($E{r}="Ramp",IF(RampWeightSum>0,$D{r}*{c}${RW}/RampWeightSum,0),$D{r}/($C{r}-$B{r}+1)),0)',
             FMT_H)
r = PR0 + NPH
put(ws, f"A{r}", "Total phase hours", font=F_BOLD, border=BOX)
for i in range(NWEEKS):
    c = wl(i)
    calc(ws, f"{c}{r}", f"=SUM({c}{PR0}:{c}{PR0 + NPH - 1})", FMT_H, bold=True)
put(ws, f"A{RW}", "Ramp weight (report development)", font=F_SUB)
for i in range(NWEEKS):
    c = wl(i)
    calc(ws, f"{c}{RW}", f"=IF(AND({c}$5>=RepStart,{c}$5<=RepEnd),MIN(1,({c}$5-RepStart+1)/(RampUp+1),(RepEnd-{c}$5+1)/(RampDown+1)),0)", "0.00")
calc(ws, f"D{RW}", f"=SUM({wl(0)}{RW}:{L(WCL)}{RW})", "0.00")
name("RampWeightSum", ws, f"D{RW}")

RH0 = RW + 3
put(ws, f"A{RH0 - 1}", "Hours by role per week", font=F_SECTION)
RC0 = RH0 + NROLES + 2
RA0 = RC0 + NROLES + 2
RF0 = RA0 + NROLES + 2
for j in range(NROLES):
    rh, rc, ra, rf = RH0 + j, RC0 + j, RA0 + j, RF0 + j
    mixcol = L(3 + j)
    for rr in (rh, rc, ra, rf):
        calc(ws, f"A{rr}", f"=Roles!B{6 + j}")
    calc(ws, f"E{rf}", f"=A{rf}")
    for i in range(NWEEKS):
        c = wl(i)
        calc(ws, f"{c}{rh}", f"=SUMPRODUCT({c}${PR0}:{c}${PR0 + NPH - 1},Roles!${mixcol}${MIX_ROW0}:${mixcol}${MIX_ROW0 + NPH - 1})", FMT_H)
        calc(ws, f"{c}{rc}", f"=IF(RoundStep>0,ROUND({c}{rh}/HrsWk/RoundStep,0)*RoundStep,{c}{rh}/HrsWk)", "0.0;-0.0;\"\"")
        inp(ws, f"{c}{ra}", None, fmt="0.0;-0.0;\"\"")
        calc(ws, f"{c}{rf}", f"=MAX(0,{c}{rc}+N({c}{ra}))", "0.0;-0.0;\"\"")
    calc(ws, f"D{rh}", f"=SUM({wl(0)}{rh}:{L(WCL)}{rh})", FMT_H)
    calc(ws, f"D{rf}", f"=SUM({wl(0)}{rf}:{L(WCL)}{rf})*HrsWk", FMT_H)
put(ws, f"A{RC0 - 1}", "Computed headcount (FTE)", font=F_SECTION)
put(ws, f"A{RA0 - 1}", "Manual adjustments (+/- FTE) - optional inputs", font=F_SECTION)
put(ws, f"A{RF0 - 1}", "Final headcount (FTE) - plan of record", font=F_SECTION)
put(ws, f"D{RF0 - 1}", "Plan hours", font=F_BOLD)
RT = RF0 + NROLES
put(ws, f"A{RT}", "Total headcount", font=F_BOLD, border=BOX)
for i in range(NWEEKS):
    c = wl(i)
    calc(ws, f"{c}{RT}", f"=SUM({c}{RF0}:{c}{RF0 + NROLES - 1})", "0.0;-0.0;\"\"", bold=True)
calc(ws, f"D{RT}", f"=SUM(D{RF0}:D{RF0 + NROLES - 1})", FMT_H, bold=True)
name("PlanHrs", ws, f"D{RT}")
name("PlanHeadcount", ws, f"{wl(0)}{RT}:{L(WCL)}{RT}")
name("PlanWeeks", ws, f"{wl(0)}5:{L(WCL)}5")
ws.freeze_panes = f"{wl(0)}7"
ws.conditional_formatting.add(f"{wl(0)}{PR0}:{L(WCL)}{PR0 + NPH - 1}",
                              CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="BDD7EE")))
ws.conditional_formatting.add(f"{wl(0)}{RF0}:{L(WCL)}{RF0 + NROLES - 1}",
                              CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="C6E0B4")))
PS0 = RT + 3
section(ws, f"A{PS0}", "Team summary", 4)
PSUM = [
    ("PeakTeam", "Peak team (FTE)", "=MAX(PlanHeadcount)", FMT_H1),
    ("AvgTeam", "Average team over the engagement (FTE)", "=IF(DurationWeeks>0,SUM(PlanHeadcount)/DurationWeeks,0)", FMT_H1),
    ("AvgDevTeam", "Average team during report iterations (FTE)", "=IF(RepWeeks>0,SUMIFS(PlanHeadcount,PlanWeeks,\">=\"&RepStart,PlanWeeks,\"<=\"&RepEnd)/RepWeeks,0)", FMT_H1),
    ("PersonWeeks", "Total person-weeks", "=SUM(PlanHeadcount)", FMT_H1),
    ("ReconVar", "Plan hours vs estimate (variance)", "=IF(TotalBaseHrs>0,PlanHrs/TotalBaseHrs-1,0)", FMT_PCT),
]
for i, (nm, lab, fml, f) in enumerate(PSUM):
    rr = PS0 + 1 + i
    put(ws, f"A{rr}", lab, border=BOX)
    calc(ws, f"D{rr}", fml, f, bold=True)
    name(nm, ws, f"D{rr}")

# ======================================================================
# Cost
# ======================================================================
ws = ws_cost
title(ws, "Cost", "Cost by role and phase, and the price build-up. Locked.")
setw(ws, {"A": 4, "B": 40, "C": 14, "D": 14, "E": 14, "F": 16, "G": 14})
section(ws, "A4", "Cost by role (estimate hours x role mix x rate card)", 7)
header_row(ws, 5, 1, ["#", "Plan role", "Location", "Hours", "Rate (USD/h)", "Cost (USD)", "Plan hours"])
for j in range(NROLES):
    r = 6 + j
    mixcol = L(3 + j)
    put(ws, f"A{r}", j + 1, border=BOX)
    calc(ws, f"B{r}", f"=Roles!B{6 + j}")
    calc(ws, f"C{r}", f"=Roles!E{6 + j}")
    calc(ws, f"D{r}", f"=SUMPRODUCT(PhaseHours,Roles!${mixcol}${MIX_ROW0}:${mixcol}${MIX_ROW0 + NPH - 1})", FMT_H)
    calc(ws, f"E{r}", f"=Roles!D{6 + j}", FMT_USD2)
    calc(ws, f"F{r}", f"=D{r}*E{r}", FMT_USD)
    calc(ws, f"G{r}", f"=Plan!D{RF0 + j}", FMT_H)
r = 6 + NROLES
put(ws, f"B{r}", "Total", font=F_BOLD, border=BOX)
for c, f in (("D", FMT_H), ("F", FMT_USD), ("G", FMT_H)):
    calc(ws, f"{c}{r}", f"=SUM({c}6:{c}{r - 1})", f, bold=True)
name("RoleHrsTotal", ws, f"D{r}")
name("RoleCostTotal", ws, f"F{r}")
name("RoleHrs", ws, f"D6:D{5 + NROLES}")
name("RoleCost", ws, f"F6:F{5 + NROLES}")
name("RoleLocCol", ws, f"C6:C{5 + NROLES}")
put(ws, f"B{r + 1}", "Implied blended rate from role mix", border=BOX)
calc(ws, f"F{r + 1}", "=IF(RoleHrsTotal>0,RoleCostTotal/RoleHrsTotal,0)", FMT_USD2)
name("ImpliedRate", ws, f"F{r + 1}")
put(ws, f"B{r + 2}", "Onsite (USA) share of hours", border=BOX)
calc(ws, f"F{r + 2}", '=IF(RoleHrsTotal>0,SUMIFS(RoleHrs,RoleLocCol,"USA")/RoleHrsTotal,0)', FMT_PCT)
name("OnsiteShare", ws, f"F{r + 2}")

C0 = r + 5
section(ws, f"A{C0}", "Cost by phase", 7)
header_row(ws, C0 + 1, 1, ["#", "Phase", "Hours", "Rate used (USD/h)", "Cost (USD)"])
for p in range(NPH):
    rr = C0 + 2 + p
    put(ws, f"A{rr}", p + 1, border=BOX)
    put(ws, f"B{rr}", PHASES[p], border=BOX)
    calc(ws, f"C{rr}", f"=INDEX(PhaseHours,{p + 1})", FMT_H)
    calc(ws, f"D{rr}", f'=IF(CostMode="Blended",BlendedRate,SUMPRODUCT(MMULT(Roles!$C${MIX_ROW0 + p}:$J${MIX_ROW0 + p},RoleRates)))', FMT_USD2)
    calc(ws, f"E{rr}", f"=C{rr}*D{rr}", FMT_USD)
rr = C0 + 2 + NPH
put(ws, f"B{rr}", "Total", font=F_BOLD, border=BOX)
calc(ws, f"C{rr}", f"=SUM(C{C0 + 2}:C{rr - 1})", FMT_H, bold=True)
calc(ws, f"E{rr}", f"=SUM(E{C0 + 2}:E{rr - 1})", FMT_USD, bold=True)
name("PhaseCostTotal", ws, f"E{rr}")
name("PhaseCost", ws, f"E{C0 + 2}:E{rr - 1}")

B0 = rr + 3
section(ws, f"A{B0}", "Price build-up", 7)
BU = [
    ("BaseCost", "1. Base cost (selected cost mode)", '=IF(CostMode="Blended",TotalBaseHrs*BlendedRate,RoleCostTotal)', FMT_USD),
    ("EffRate", "   Effective rate (USD/h)", "=IF(TotalBaseHrs>0,BaseCost/TotalBaseHrs,0)", FMT_USD2),
    ("ContCost", "2. Contingency (by input confidence)", "=ContHrs*EffRate", FMT_USD),
    ("RiskCost", "3. Risk reserve (Risks tab)", "=RiskHrs*EffRate", FMT_USD),
    ("ServicesGross", "   Services before discount", "=BaseCost+ContCost+RiskCost", FMT_USD),
    ("DiscountAmt", "   Discount on bill rates", "=-ServicesGross*DiscountPct", FMT_USD),
    ("ServicesNet", "   Services price", "=ServicesGross+DiscountAmt", FMT_USD),
    ("PassThrough", "4. Pass-throughs (tools, travel, other)", "=PassTools+PassTravel+PassOther", FMT_USD),
    ("FixedPrice", "5. FIXED PRICE", "=ServicesNet+PassThrough", FMT_USD),
]
for i, (nm, lab, fml, f) in enumerate(BU):
    rr = B0 + 1 + i
    bold = nm == "FixedPrice"
    put(ws, f"B{rr}", lab, border=BOX, font=F_BOLD if bold else F_BASE)
    calc(ws, f"F{rr}", fml, f, bold=bold, fill=FILL_SECTION if bold else None)
    name(nm, ws, f"F{rr}")
M1 = B0 + len(BU) + 3
section(ws, f"A{M1}", "Payment milestones", 7)
header_row(ws, M1 + 1, 1, ["#", "Milestone", "%", "Amount (USD)"])
for i in range(6):
    rr = M1 + 2 + i
    put(ws, f"A{rr}", i + 1, border=BOX)
    calc(ws, f"B{rr}", f'=IF(INDEX(MilestoneNames,{i + 1})="","",INDEX(MilestoneNames,{i + 1}))')
    calc(ws, f"C{rr}", f"=INDEX(MilestonePct,{i + 1})", FMT_PCT0)
    calc(ws, f"D{rr}", f"=FixedPrice*C{rr}", FMT_USD)

# ======================================================================
# Checks
# ======================================================================
ws = ws_chk
title(ws, "Checks", "Every check must PASS before the estimate is quoted. The Summary banner reads from here.")
setw(ws, {"A": 4, "B": 62, "C": 10, "D": 70})
header_row(ws, 4, 1, ["#", "Check", "Result", "Detail"])
ERR_RANGES = ["Effort!A7:AV60", "Schedule!B5:H80", f"Plan!A5:{L(WCL)}{PS0 + 6}", "Cost!A5:G60"]
CHECKS = [
    ("At least one active migration pair", "ActivePairs>=1", '=ActivePairs&" active pair(s)"'),
    ("Every active pair maps to source and target archetypes", "FlagMap=0", '=FlagMap&" pair(s) with an unmapped tool - add the tool on Lists"'),
    ("Report and semantic-model complexity mixes total 100%", "FlagMix=0", '=FlagMix&" pair(s) with a mix that does not total 100%"'),
    ("Retired + Consolidated <= 100% per pair", "FlagRat=0", '=FlagRat&" pair(s) over 100%"'),
    ("Every pair-factor override has a reason", "FlagWhy=0", '=FlagWhy&" override(s) without a reason"'),
    ("Pages per report >= 1 and at least one target model per pair", "FlagPpr=0", '=FlagPpr&" pair(s) to fix"'),
    ("Role mix rows each total 100%", f"SUMPRODUCT(--(ABS(RoleMixTotals-1)>0.0001))=0", '=SUMPRODUCT(--(ABS(RoleMixTotals-1)>0.0001))&" phase row(s) off"'),
    ("Every plan role has a rate", "COUNTIF(RoleRates,\">0\")=COUNTA(RoleNames)", '=COUNTIF(RoleRates,">0")&" of "&COUNTA(RoleNames)&" roles priced"'),
    ("Payment milestones total 100%", "ABS(MilestoneTotal-1)<0.0001", '=TEXT(MilestoneTotal,"0%")'),
    ("Plan reconciles to estimate within tolerance", "ABS(ReconVar)<=ReconTol", '="Variance "&TEXT(ReconVar,"0.0%")&" (tolerance "&TEXT(ReconTol,"0%")&")"'),
    (f"Duration fits the {NWEEKS}-week plan horizon", f"DurationWeeks<={NWEEKS}", '=DurationWeeks&" weeks"'),
    ("Totals tie: phases = pairs + shared", "ABS(TotalBaseHrs-PairTotalHrs)<1", '="Phases "&TEXT(TotalBaseHrs,"#,##0")&" h vs pairs "&TEXT(PairTotalHrs,"#,##0")&" h"'),
    ("Totals tie: role hours = phase hours", "ABS(RoleHrsTotal-TotalBaseHrs)<1", '="Roles "&TEXT(RoleHrsTotal,"#,##0")&" h"'),
    ("Totals tie: cost by phase = cost by role (role-mix mode)", 'OR(CostMode="Blended",ABS(PhaseCostTotal-RoleCostTotal)<1)', '="Phase cost "&TEXT(PhaseCostTotal,"$#,##0")'),
    ("No error values in calculation tabs", "+".join(f"SUMPRODUCT(--ISERROR({r}))" for r in ERR_RANGES) + "=0",
     "=" + "+".join(f"SUMPRODUCT(--ISERROR({r}))" for r in ERR_RANGES) + '&" error cell(s)"'),
]
for i, (lab, cond, detail) in enumerate(CHECKS):
    r = 5 + i
    put(ws, f"A{r}", i + 1, border=BOX)
    put(ws, f"B{r}", lab, border=BOX)
    calc(ws, f"C{r}", f'=IF(IFERROR({cond},FALSE),"PASS","FAIL")', bold=True)
    calc(ws, f"D{r}", f'=IFERROR({detail[1:]},"")')
CK1 = 5 + len(CHECKS) - 1
ws.conditional_formatting.add(f"C5:C{CK1}", CellIsRule(operator="equal", formula=['"PASS"'], fill=FILL_GOOD))
ws.conditional_formatting.add(f"C5:C{CK1}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FILL_BAD))
r = CK1 + 2
put(ws, f"B{r}", "Failed checks", font=F_BOLD, border=BOX)
calc(ws, f"C{r}", f'=COUNTIF(C5:C{CK1},"FAIL")', FMT_INT, bold=True)
name("FailCount", ws, f"C{r}")

# ======================================================================
# Summary
# ======================================================================
ws = ws_sum
ws.sheet_view.showGridLines = False
setw(ws, {"A": 2, "B": 34, "C": 24, "D": 16, "E": 16, "F": 18, "G": 3, "H": 30, "I": 16, "J": 16, "K": 16, "L": 16})
put(ws, "B1", '="BI Migration Estimate - "&Client', font=F_TITLE)
put(ws, "B2", '=Opportunity&"  |  "&EstVersion&"  |  "&TEXT(EstDate,"dd-mmm-yyyy")&"  |  Start "&TEXT(StartDate,"dd-mmm-yyyy")&"  |  All money in "&Currency', font=F_SUB)
ws.merge_cells("B4:L4")
put(ws, "B4", '=IF(FailCount=0,"VALID ESTIMATE - all checks pass","NOT VALID - "&FailCount&" check(s) failing: see the Checks tab")',
    font=Font(name=FONT, size=12, bold=True), align=CENTER)
ws.row_dimensions[4].height = 24
ws.conditional_formatting.add("B4:L4", FormulaRule(formula=["FailCount=0"], fill=FILL_GOOD))
ws.conditional_formatting.add("B4:L4", FormulaRule(formula=["FailCount>0"], fill=FILL_BAD))
KPIS = [
    ("Fixed price", "=FixedPrice", FMT_USD), ("Total effort (hours)", "=TotalHrs", FMT_H),
    ("Duration (working weeks)", "=DurationWeeks", FMT_INT), ("Planned end date", "=EndDate", FMT_DATE),
    ("Development iterations", '=IterCount&" x "&IterLen&" wk"', None), ("Releases", "=Releases", FMT_INT),
    ("Peak team (FTE)", "=PeakTeam", FMT_H1), ("Average team (FTE)", "=AvgTeam", FMT_H1),
]
SPANS = [("B", "C"), ("D", "F"), ("H", "I"), ("J", "L")]
positions = [(a, b, r) for r in (6, 9) for a, b in SPANS]
for (lab, fml, f), (c, col2, r) in zip(KPIS, positions):
    ws.merge_cells(f"{c}{r}:{col2}{r}")
    ws.merge_cells(f"{c}{r + 1}:{col2}{r + 1}")
    put(ws, f"{c}{r}", lab, font=F_KPI_LABEL, fill=FILL_KPI, align=CENTER)
    put(ws, f"{c}{r + 1}", fml, font=F_KPI, fmt=f, fill=FILL_KPI, align=CENTER)
    ws.row_dimensions[r + 1].height = 30
put(ws, "B12", '="Average team during report iterations: "&TEXT(AvgDevTeam,"0.0")&" FTE  |  Contingency "&TEXT(ContPct,"0.0%")&"  |  Onsite share "&TEXT(OnsiteShare,"0.0%")&"  |  Cost mode: "&CostMode', font=F_SUB)

section(ws, "B14", "Effort and cost by phase", 5)
header_row(ws, 15, 2, ["Phase", "Hours", "% of hours", "Cost (USD)", "Weeks (start-end)"])
for p in range(NPH):
    r = 16 + p
    put(ws, f"B{r}", PHASES[p], border=BOX)
    calc(ws, f"C{r}", f"=INDEX(PhaseHours,{p + 1})", FMT_H)
    calc(ws, f"D{r}", f"=IF(TotalBaseHrs>0,C{r}/TotalBaseHrs,0)", FMT_PCT)
    calc(ws, f"E{r}", f"=INDEX(PhaseCost,{p + 1})", FMT_USD)
    calc(ws, f"F{r}", f'=IF(C{r}>0,INDEX(PhStart,{p + 1})&" - "&INDEX(PhEnd,{p + 1}),"")')
r = 16 + NPH
SUMROWS = [("Base effort", "=TotalBaseHrs", "=BaseCost"), ("Contingency", "=ContHrs", "=ContCost"),
           ("Risk reserve", "=RiskHrs", "=RiskCost"), ("Discount", "", "=DiscountAmt"),
           ("Pass-throughs", "", "=PassThrough"), ("FIXED PRICE", "=TotalHrs", "=FixedPrice")]
for i, (lab, h, c) in enumerate(SUMROWS):
    rr = r + i
    bold = lab == "FIXED PRICE"
    put(ws, f"B{rr}", lab, font=F_BOLD, border=BOX)
    calc(ws, f"C{rr}", h if h else None, FMT_H, bold=bold)
    calc(ws, f"E{rr}", c, FMT_USD, bold=bold, fill=FILL_SECTION if bold else None)

section(ws, "H14", "Team and cost by role", 5)
header_row(ws, 15, 8, ["Role", "Location", "Hours", "Rate (USD/h)", "Cost (USD)"])
for j in range(NROLES):
    r = 16 + j
    calc(ws, f"H{r}", f"=Cost!B{6 + j}")
    calc(ws, f"I{r}", f"=Cost!C{6 + j}")
    calc(ws, f"J{r}", f"=Cost!D{6 + j}", FMT_H)
    calc(ws, f"K{r}", f"=Cost!E{6 + j}", FMT_USD2)
    calc(ws, f"L{r}", f"=Cost!F{6 + j}", FMT_USD)
r = 16 + NROLES
put(ws, f"H{r}", "Total", font=F_BOLD, border=BOX)
calc(ws, f"J{r}", "=RoleHrsTotal", FMT_H, bold=True)
calc(ws, f"L{r}", "=RoleCostTotal", FMT_USD, bold=True)

section(ws, "H27", "Key dates", 5)
header_row(ws, 28, 8, ["Milestone", "Week", "Date"])
for i in range(MIL_N):
    rr = 29 + i
    calc(ws, f"H{rr}", f"=INDEX(MilestoneLabels,{i + 1})")
    calc(ws, f"I{rr}", f'=INDEX(MilestoneWeeks,{i + 1})', FMT_INT)
    calc(ws, f"J{rr}", f'=INDEX(MilestoneDates,{i + 1})', FMT_DATE)

PT0 = 36
section(ws, f"B{PT0}", "Effort by migration pair", 5)
header_row(ws, PT0 + 1, 2, ["Pair", "Source -> target", "Reports to build", "Hours (incl. shared)", "Contingency %"])
for p in range(NPAIRS):
    r = PT0 + 2 + p
    er = SR0 + p
    calc(ws, f"B{r}", f'=IF(Effort!C{er}=1,Effort!B{er},"")')
    calc(ws, f"C{r}", f'=IF(Effort!C{er}=1,{S("src", er)}&" -> "&{S("tgt", er)},"")')
    calc(ws, f"D{r}", f'=IF(Effort!C{er}=1,Effort!R{er},"")', FMT_H)
    calc(ws, f"E{r}", f'=IF(Effort!C{er}=1,Effort!AP{er},"")', FMT_H)
    calc(ws, f"F{r}", f'=IF(Effort!C{er}=1,Effort!AN{er},"")', FMT_PCT0)

chart = BarChart()
chart.type = "col"
chart.grouping = "stacked"
chart.overlap = 100
chart.title = "Weekly headcount by role (FTE)"
chart.y_axis.title = "FTE"
chart.x_axis.title = "Working week"
data = Reference(ws_plan, min_col=5, max_col=WCL, min_row=RF0, max_row=RF0 + NROLES - 1)
cats = Reference(ws_plan, min_col=WC0, max_col=WCL, min_row=5, max_row=5)
chart.add_data(data, from_rows=True, titles_from_data=True)
chart.set_categories(cats)
chart.height = 9
chart.width = 26
chart.gapWidth = 20
ws.add_chart(chart, f"H{29 + MIL_N + 2}")

# ======================================================================
# Assumptions (printable)
# ======================================================================
ws = ws_asm
title(ws, "Assumptions & exclusions", "Attach to the proposal. Every value is live from the input and reference tabs.")
setw(ws, {"A": 4, "B": 58, "C": 28, "D": 50})
ar = [4]


def aline(label, fml, fmt=None, note=None):
    r = ar[0]
    put(ws, f"B{r}", label, border=BOX)
    calc(ws, f"C{r}", fml, fmt)
    if note:
        put(ws, f"D{r}", note, border=BOX, align=WRAP)
    ar[0] += 1


def asec(t):
    ar[0] += 1
    section(ws, f"B{ar[0]}", t, 3)
    ar[0] += 1


asec("Engagement")
aline("Client / opportunity", '=Client&" - "&Opportunity')
aline("Estimate version and date", '=EstVersion&" - "&TEXT(EstDate,"dd-mmm-yyyy")')
aline("Planned start", "=StartDate", FMT_DATE)
aline("Engagement type", "=EngType", None, "Bundled engagements carry PM in the ETL/EDW programme")
asec("Delivery assumptions")
aline("Productive hours per person-week", "=HrsWk", FMT_H)
aline("Iteration length (weeks)", "=IterLen", FMT_INT)
aline("Development iterations", "=IterCount", FMT_INT)
aline("Solve mode", "=SolveMode")
aline("Peak report-development team (FTE)", "=PeakRepTeam", FMT_H1)
aline("Ramp-up / ramp-down weeks", '=RampUp&" / "&RampDown')
aline("Model lead time before report iterations (weeks)", "=ModelLead", FMT_INT)
aline("SIT: weeks per iteration / engineers / fixed h per iteration", '=SITWk&" / "&SITEng&" / "&SITFixed')
aline("UAT support: weeks per iteration / engineers", '=UATWk&" / "&UATEng', None, "1 engineer up to the volume threshold, 2 above it")
aline("Deployment support: weeks per iteration / engineers", '=DepWk&" / "&DepEng')
aline("Deployment pipeline available", "=Pipeline")
aline("Hypercare weeks per go-live / crew", '=HCWk&" / "&HCCrew')
aline("Holiday allowance (calendar dates)", "=HolidayPct", FMT_PCT0)
asec("Automation (global % x automatable cap = saving)")
for lab, nm, cap in AUTO:
    aline(lab, f'=TEXT({nm},"0%")&" x "&TEXT({cap},"0%")&" = "&TEXT({nm}*{cap},"0.0%")')
asec("Estimation basis")
aline("Hours per report page (S / M / C / VC / Derived)", '=INDEX(HrsPage,1)&" / "&INDEX(HrsPage,2)&" / "&INDEX(HrsPage,3)&" / "&INDEX(HrsPage,4)&" / "&HrsDerived', None, "Carried from the previous calculator")
aline("Learning gain by quarter of iterations", '=TEXT(INDEX(LearnQ,1),"0%")&" / "&TEXT(INDEX(LearnQ,2),"0%")&" / "&TEXT(INDEX(LearnQ,3),"0%")&" / "&TEXT(INDEX(LearnQ,4),"0%")')
aline("Semantic model hours (S / M / C / VC)", '=INDEX(ModelHrs,1)&" / "&INDEX(ModelHrs,2)&" / "&INDEX(ModelHrs,3)&" / "&INDEX(ModelHrs,4)')
aline("Contingency (weighted by input confidence)", "=ContPct", FMT_PCT)
aline("Risk reserve hours", "=RiskHrs", FMT_H)
aline("Cost mode / discount", '=CostMode&" / "&TEXT(DiscountPct,"0%")')
asec("Pair overrides (factor overrides must carry a reason)")
header_row(ws, ar[0], 2, ["Pair", "Suggested -> used factor", "Reason"])
ar[0] += 1
for p in range(NPAIRS):
    er = SR0 + p
    r = ar[0]
    calc(ws, f"B{r}", f'=IF(Effort!C{er}=1,Effort!B{er},"")')
    calc(ws, f"C{r}", f'=IF(Effort!C{er}=1,TEXT(Effort!F{er},"0.00")&" -> "&TEXT(Effort!G{er},"0.00"),"")')
    calc(ws, f"D{r}", f'=IF(Effort!C{er}=1,IF({S("why", er)}="","",{S("why", er)}),"")')
    ar[0] += 1
asec("Exclusions")
for t in ["Data / ETL / EDW migration and data reconciliation of the underlying platform (estimated separately).",
          "BI and data-platform licences and infrastructure, unless entered as a pass-through.",
          "Remediation of source data quality issues.",
          "Reports retired or consolidated by rationalisation (only merge effort is included).",
          "End-user training beyond the training and handover allowance per pair.",
          "Work outside the scope rows and risk register in this workbook."]:
    put(ws, f"B{ar[0]}", t, align=WRAP)
    ws.merge_cells(f"B{ar[0]}:D{ar[0]}")
    ar[0] += 1

# ======================================================================
# Scenarios
# ======================================================================
ws = ws_scn
title(ws, "Scenarios", "Compare options: copy the Current column and Paste Special > Values into Scenario A, B or C, then change inputs.")
setw(ws, {"A": 2, "B": 40, "C": 18, "D": 18, "E": 18, "F": 18})
header_row(ws, 4, 2, ["Measure", "Current", "Scenario A", "Scenario B", "Scenario C"])
put(ws, "B5", "Scenario name", border=BOX)
calc(ws, "C5", '="Current"')
for c in "DEF":
    inp(ws, f"{c}5", None)
SCN = [("Fixed price (USD)", "=FixedPrice", FMT_USD), ("Total hours", "=TotalHrs", FMT_H),
       ("Base hours", "=TotalBaseHrs", FMT_H), ("Duration (weeks)", "=DurationWeeks", FMT_INT),
       ("Development iterations", "=IterCount", FMT_INT), ("Peak team (FTE)", "=PeakTeam", FMT_H1),
       ("Average team (FTE)", "=AvgTeam", FMT_H1), ("Report automation", "=AutoRep", FMT_PCT0),
       ("Contingency %", "=ContPct", FMT_PCT), ("Cost mode", "=CostMode", None)]
for i, (lab, fml, f) in enumerate(SCN):
    r = 6 + i
    put(ws, f"B{r}", lab, border=BOX)
    calc(ws, f"C{r}", fml, f)
    for c in "DEF":
        inp(ws, f"{c}{r}", None, fmt=f)

# ======================================================================
# ChangeLog
# ======================================================================
ws = ws_log
title(ws, "Change log", "Record every change to inputs, baseline values or formulas.")
setw(ws, {"A": 12, "B": 14, "C": 20, "D": 80})
header_row(ws, 4, 1, ["Version", "Date", "Author", "Change"])
inp(ws, "A5", "v2.0")
inp(ws, "B5", dt.date(2026, 10, 13), fmt=FMT_DATE)
inp(ws, "C5", "BI Estimator build")
inp(ws, "D5", "First release of v2, built from the approved requirements document.")
for r in range(6, 30):
    for c in "ABCD":
        inp(ws, f"{c}{r}", None, fmt=FMT_DATE if c == "B" else None)

# ======================================================================
# Guide
# ======================================================================
ws = ws_guide
title(ws, "BI Migration Estimator v2", "Quotable fixed-price estimate for BI report and semantic-model migrations.")
setw(ws, {"A": 2, "B": 24, "C": 100})
put(ws, "B4", "How to use", font=F_SECTION)
STEPS = [
    ("1. Setup", "Enter engagement details, iteration length, team and ramp settings, SIT/UAT/deployment cadence, automation % by phase and cost mode."),
    ("2. Scope", "One row per migration pair. Pick or type the source and target BI tool and data platform, then inventory, complexity mix, rationalisation, semantic models, security and AI readiness. Example rows are provided - overwrite them."),
    ("3. Risks", "List named risks with impact hours and probability. Expected hours become the risk reserve."),
    ("4. Roles", "Choose the rate-card row for each plan role and adjust the role mix per phase (each row totals 100%)."),
    ("5. Checks", "Fix every FAIL. The Summary banner turns green only when all checks pass."),
    ("6. Summary", "Read the price, effort, duration, team and dates. Print Summary and Assumptions for the proposal."),
    ("7. Plan", "Review the weekly headcount. Use the manual adjustment rows to reshape the team; the reconciliation check flags gaps over the tolerance."),
]
for i, (a, b) in enumerate(STEPS):
    put(ws, f"B{5 + i}", a, font=F_BOLD)
    put(ws, f"C{5 + i}", b, align=WRAP)
put(ws, "B13", "Colour key", font=F_SECTION)
put(ws, "B14", "Input", font=F_INPUT, fill=FILL_INPUT, border=BOX)
put(ws, "C14", "Yellow cell with blue text: editable input (all other cells are locked; sheets are protected without a password).")
put(ws, "B15", "Calculated", border=BOX)
put(ws, "C15", "Black text: formula - do not overwrite.")
put(ws, "B16", "Header", font=F_HEAD, fill=FILL_HEAD)
put(ws, "C16", "Table header.")
put(ws, "B18", "Tabs", font=F_SECTION)
TABS = [
    ("Summary", "One-page quotable view."), ("Setup", "Engagement-wide inputs."), ("Scope", "Migration pairs (up to 20)."),
    ("Risks", "Risk register."), ("Roles", "Plan roles, rate selection and role mix by phase."),
    ("Rates", "Rate card (bill rates, USD/h)."), ("Lists", "Tools, archetypes, known pair factors, data platforms."),
    ("Baseline", "Reference values with status: Carried / Reshaped / New proposal / User decision."),
    ("Effort", "Hours per pair and phase; programme-level hours; totals."),
    ("Schedule", "Iterations, waves, phase windows, milestones."), ("Plan", "Weekly hours and headcount by role."),
    ("Cost", "Cost by role and phase; price build-up; payment milestones."),
    ("Assumptions", "Printable assumptions and exclusions."), ("Checks", "Validation; must all PASS."),
    ("Scenarios", "Side-by-side comparison of saved options."), ("ChangeLog", "Version history."),
]
for i, (a, b) in enumerate(TABS):
    put(ws, f"B{19 + i}", a, font=F_BOLD)
    put(ws, f"C{19 + i}", b)
r = 19 + len(TABS) + 1
put(ws, f"B{r}", "Method", font=F_SECTION)
METHOD = [
    "Report development = full-effort pages x mix-weighted hours per page x pair factor x data-platform uplift x report-type factor x (1 - learning gain) x (1 - automation saving), plus derived and merge hours.",
    "Semantic models = target models x mix-weighted model hours x approach share x pair factor x platform uplift, plus consolidation hours, less automation saving.",
    "Iterations: report hours / (peak team x hours per week) plus half the ramp weeks, divided by iteration length (or fixed directly).",
    "SIT runs one iteration behind development and UAT two behind; both are sized per iteration (crew x weeks), plus a final regression and UAT cycle.",
    "Automation saving per phase = automation % x automatable cap, so automation never removes the review and validation work.",
    "No step tiers: initiation and analysis interpolate between the old calculator's tier values.",
]
for i, t in enumerate(METHOD):
    put(ws, f"C{r + 1 + i}", t, align=WRAP)

# ======================================================================
# Fonts everywhere + protection
# ======================================================================
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, size=f.size or 10, bold=f.bold, italic=f.italic, color=f.color)
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False

for ws in (ws_sum, ws_asm):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.active = wb.sheetnames.index("Summary")
wb.save(OUT)
print(f"wrote {OUT} ({len(NAMES)} names)")
